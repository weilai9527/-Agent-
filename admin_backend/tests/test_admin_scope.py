from __future__ import annotations

import os
from pathlib import Path
import subprocess
import sys
import textwrap

from cryptography.fernet import Fernet


def test_scoped_student_access_and_fresh_admin_database(tmp_path):
    project_dir = Path(__file__).resolve().parents[2]
    environment = os.environ.copy()
    environment.update(
        {
            "APP_ENV": "production",
            "DB_ENGINE": "sqlite",
            "SQLITE_PATH": str(tmp_path / "admin-scope.sqlite"),
            "ADMIN_MASTER_EMAIL": "scope-root@example.com",
            "ADMIN_MASTER_PASSWORD": "ScopeRootPassword123!",
            "ADMIN_COOKIE_SECURE": "false",
            "STUDENT_TEMP_PASSWORD_KEY": Fernet.generate_key().decode("ascii"),
        }
    )
    script = textwrap.dedent(
        """
        import json
        from io import BytesIO
        from uuid import uuid4

        from fastapi.testclient import TestClient
        from openpyxl import load_workbook

        import admin_backend.src.main as admin_main
        from admin_backend.src.security import encrypt_temporary_password, hash_password

        origin = {"Origin": "http://127.0.0.1:5174"}
        root = TestClient(admin_main.app)
        login = root.post(
            "/api/admin/auth/login",
            headers=origin,
            json={"email": "scope-root@example.com", "password": "ScopeRootPassword123!"},
        )
        assert login.status_code == 200, login.text
        # Importing only the management service must initialize the shared tables.
        assert root.get("/api/admin/snapshot").status_code == 200

        college = root.post(
            "/api/admin/campus/colleges", headers=origin,
            json={"code": "SCOPE-COLLEGE", "name": "Scope college"},
        )
        assert college.status_code == 201, college.text
        college_id = college.json()["college"]["id"]
        program = root.post(
            "/api/admin/campus/programs", headers=origin,
            json={"collegeId": college_id, "name": "Scope program"},
        )
        assert program.status_code == 201, program.text
        program_id = program.json()["program"]["id"]
        classes = []
        for name in ("Allowed", "Outside"):
            response = root.post(
                "/api/admin/campus/classes", headers=origin,
                json={"programId": program_id, "name": name, "graduationYear": 2027},
            )
            assert response.status_code == 201, response.text
            classes.append(response.json()["class"]["id"])

        for user_id, student_no, class_id in (
            ("allowed-student", "20260001", classes[0]),
            ("outside-student", "20260002", classes[1]),
        ):
            password = "TemporaryStudentPassword123!"
            admin_main.db.execute(
                "INSERT INTO users (id, email, password_hash, name, student_no, "
                "must_change_password, temp_password_encrypted) VALUES (?, ?, ?, ?, ?, 1, ?)",
                (user_id, f"{student_no}@student.local", hash_password(password),
                 student_no, student_no, encrypt_temporary_password(password)),
            )
            admin_main.upsert_student_enrollment(user_id, class_id=class_id, student_no=student_no)

        admin_main.db.execute(
            "INSERT INTO admin_users "
            "(id, email, password_hash, name, role, status, permissions, student_scope) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            ("limited-admin", "limited@example.com", hash_password("LimitedPassword123!"),
             "Limited", "manager", "normal", json.dumps(["manageStudents"]),
             json.dumps([{"college": college_id, "program": program_id, "class": classes[0]}])),
        )
        admin_main.db.commit()
        limited = TestClient(admin_main.app)
        login = limited.post(
            "/api/admin/auth/login", headers=origin,
            json={"email": "limited@example.com", "password": "LimitedPassword123!"},
        )
        assert login.status_code == 200, login.text
        snapshot = limited.get("/api/admin/snapshot")
        assert snapshot.status_code == 200, snapshot.text
        assert {item["id"] for item in snapshot.json()["candidates"]} == {"allowed-student"}
        assert snapshot.json()["permissions"]["canViewConnectionLogs"] is False
        assert snapshot.json()["permissions"]["canImportStudentAccounts"] is False

        assert limited.get("/api/admin/candidates/allowed-student").status_code == 200
        assert limited.get("/api/admin/candidates/outside-student").status_code == 403
        assert limited.get("/api/admin/student-accounts/allowed-student/temporary-password").status_code == 200
        assert limited.get("/api/admin/student-accounts/outside-student/temporary-password").status_code == 403
        before = admin_main.one("SELECT password_hash FROM users WHERE id = ?", ("outside-student",))
        reset = limited.post(
            "/api/admin/student-accounts/outside-student/reset-password", headers=origin,
        )
        assert reset.status_code == 403, reset.text
        after = admin_main.one("SELECT password_hash FROM users WHERE id = ?", ("outside-student",))
        assert before == after
        export = limited.get("/api/admin/student-accounts/export")
        assert export.status_code == 200, export.text
        rows = list(load_workbook(BytesIO(export.content), data_only=True).active.iter_rows(values_only=True))
        assert len(rows) == 2
        assert rows[1][1] == "20260001"
        assert limited.get("/api/admin/connection-logs").status_code == 403
        assert limited.post(
            "/api/admin/student-accounts/import", headers=origin,
            files={"file": ("students.xlsx", b"invalid", "application/octet-stream")},
        ).status_code == 403
        admin_main.db.execute(
            "INSERT INTO admin_users "
            "(id, email, password_hash, name, role, status, permissions, student_scope) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            ("empty-admin", "empty@example.com", hash_password("EmptyPassword123!"),
             "Empty", "manager", "normal", json.dumps(["manageStudents"]), json.dumps([])),
        )
        admin_main.db.commit()
        empty = TestClient(admin_main.app)
        assert empty.post(
            "/api/admin/auth/login", headers=origin,
            json={"email": "empty@example.com", "password": "EmptyPassword123!"},
        ).status_code == 200
        assert empty.get("/api/admin/candidates/allowed-student").status_code == 403
        empty_export = empty.get("/api/admin/student-accounts/export")
        assert empty_export.status_code == 200
        assert load_workbook(BytesIO(empty_export.content)).active.max_row == 1
        listing = root.get("/api/admin/student-registrations")
        assert listing.status_code == 200, listing.text
        assert listing.json()["registrations"] == []
        imported = root.post(
            "/api/admin/student-registrations/import", headers=origin,
            files={
                "file": (
                    "students.csv",
                    b"student_no,name,college,gender,class_name,counselor\\n"
                    b"20260003,Scope New,Scope college,female,2301CP,Ms. Wang\\n",
                    "text/csv",
                )
            },
        )
        assert imported.status_code == 200, imported.text
        assert imported.json()["imported"] == 1
        assert imported.json()["activated"] == 1
        listing = root.get("/api/admin/student-registrations")
        assert listing.status_code == 200, listing.text
        assert listing.json()["registrations"][0]["studentNo"] == "20260003"
        assert listing.json()["registrations"][0]["activated"] is True
        assert root.post(
            "/api/admin/organization/import", headers=origin,
            files={"file": ("old.csv", b"student_no,name", "text/csv")},
        ).status_code == 410
        """
    )
    completed = subprocess.run(
        [sys.executable, "-c", script],
        cwd=project_dir,
        env=environment,
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr
