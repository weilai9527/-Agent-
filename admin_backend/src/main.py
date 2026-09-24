from __future__ import annotations

import csv
import io
import json
import logging
import os
import re
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .database import DB_ENGINE, all_rows, db, ensure_admin_schema, get_database_path, one
from .security import (
    create_token,
    decrypt_temporary_password,
    encrypt_temporary_password,
    generate_temporary_password,
    hash_password,
    hash_token,
    is_valid_email,
    normalize_email,
    sanitize_admin,
    verify_password,
)
from shared.career_catalog import (
    approve_job_suggestion_to_draft,
    catalog_tree,
    create_catalog_entity,
    create_version,
    delete_catalog_entity,
    has_permission,
    import_catalog_excel,
    list_job_suggestions,
    list_versions,
    merge_job_suggestion,
    publish_version,
    replace_job_competencies,
    review_job_suggestion,
    update_catalog_entity,
)
from shared.recruitment import (
    create_jd_submission,
    create_job_posting,
    delete_jd_submission,
    delete_job_posting,
    get_jd_submission,
    get_job_match,
    get_job_posting,
    list_jd_submissions,
    list_job_matches,
    list_job_postings,
    mark_jd_submission_approved,
    update_jd_submission,
    update_job_posting,
    upsert_job_posting_by_source,
)


app = FastAPI(title="Multi Agent Interview Admin API")
logger = logging.getLogger(__name__)
PROJECT_DIR = Path(__file__).resolve().parents[2]
USER_BACKEND_ENV_PATH = PROJECT_DIR / "backend" / ".env"
ADMIN_SESSION_COOKIE = "admin_session"
ADMIN_SESSION_MAX_AGE = int(os.environ.get("ADMIN_SESSION_MAX_AGE", str(60 * 60 * 8)))
# 超级管理员账号定义在代码中（可用环境变量覆盖），不存入数据库由界面创建/删除。
MASTER_ADMIN_EMAIL = normalize_email(
    os.environ.get("ADMIN_MASTER_EMAIL")
    or os.environ.get("ADMIN_BOOTSTRAP_EMAIL")
    or "admin@ai.local"
)
MASTER_ADMIN_PASSWORD = (
    os.environ.get("ADMIN_MASTER_PASSWORD")
    or os.environ.get("ADMIN_BOOTSTRAP_PASSWORD")
    or "Admin@2026!Master"
)
MASTER_ADMIN_NAME = (
    os.environ.get("ADMIN_MASTER_NAME")
    or os.environ.get("ADMIN_BOOTSTRAP_NAME")
    or "超级管理员"
).strip()[:80] or "超级管理员"
# 固定权限点集合：下级管理员的权限通过这些权限点授予，角色名仅作展示标签。
ALL_PERMISSIONS = (
    "manageCatalog",
    "manageStudents",
    "manageOrganization",
    "manageModelConfig",
    "viewInterviews",
    "viewReports",
    "viewConnectionLogs",
)
LOGIN_WINDOW_SECONDS = 15 * 60
MAX_LOGIN_ATTEMPTS = 8
login_attempts: dict[str, dict[str, int | float]] = {}
WEBRTC_DIAGNOSTIC_RETENTION_DAYS = max(
    1,
    min(365, int(os.environ.get("WEBRTC_DIAGNOSTIC_RETENTION_DAYS", "14"))),
)

allowed_origins = [
    origin.strip()
    for origin in os.environ.get(
        "ADMIN_FRONTEND_ORIGIN",
        "http://127.0.0.1:5174,http://127.0.0.1:5175,http://127.0.0.1:5176,"
        "http://localhost:5174,http://localhost:5175,http://localhost:5176",
    ).split(",")
    if origin.strip()
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type"],
)


def error(status_code: int, message: str) -> HTTPException:
    return HTTPException(status_code=status_code, detail={"error": message})


@app.exception_handler(HTTPException)
async def http_exception_handler(_request: Request, exc: HTTPException):
    detail = exc.detail
    if isinstance(detail, dict) and "error" in detail:
        return JSONResponse(status_code=exc.status_code, content={"error": detail["error"]})
    return JSONResponse(status_code=exc.status_code, content={"error": str(detail)})


def utc_after(seconds: int) -> str:
    return (datetime.now(timezone.utc) + timedelta(seconds=seconds)).strftime("%Y-%m-%d %H:%M:%S")


def request_ip(request: Request) -> str:
    return request.client.host if request.client else "unknown"


def validate_admin_origin(request: Request) -> None:
    origin = (request.headers.get("origin") or "").strip()
    if origin and origin not in allowed_origins:
        raise error(403, "管理端请求来源不受信任。")


def find_admin_by_session(request: Request) -> dict | None:
    token = request.cookies.get(ADMIN_SESSION_COOKIE)
    if not token:
        return None
    return one(
        """
        SELECT users.id, users.email, users.name, users.role, users.status,
               users.permissions, users.student_scope, users.last_login_at
        FROM admin_sessions AS sessions
        JOIN admin_users AS users ON users.id = sessions.admin_user_id
        WHERE sessions.token_hash = ?
          AND sessions.expires_at > ?
          AND users.status = 'normal'
        """,
        (hash_token(token), utc_after(0)),
    )


def require_admin(request: Request) -> dict:
    admin = find_admin_by_session(request)
    if not admin:
        raise error(401, "管理员登录已失效，请重新登录。")
    return admin


def is_super_admin(admin: dict) -> bool:
    """只有代码中定义的主管理员账号是超级管理员，避免自由文本角色越权。"""
    return str(admin.get("email") or "").lower() == MASTER_ADMIN_EMAIL


def admin_permissions(admin: dict) -> set[str]:
    raw = admin.get("permissions")
    if isinstance(raw, list):
        return {str(item) for item in raw if item}
    if isinstance(raw, str):
        try:
            data = json.loads(raw)
        except (TypeError, ValueError):
            return set()
        return {str(item) for item in data if item} if isinstance(data, list) else set()
    return set()


def admin_student_scope(admin: dict) -> list[dict[str, str]] | None:
    """返回管理员的学生数据范围；None 表示不受限（超级管理员），[] 表示无任何范围。"""
    if is_super_admin(admin):
        return None
    raw = admin.get("student_scope")
    if isinstance(raw, list):
        return raw
    if isinstance(raw, str):
        try:
            data = json.loads(raw)
        except (TypeError, ValueError):
            return []
        return data if isinstance(data, list) else []
    return []


def scope_allows(scope: list[dict[str, str]] | None, college_id: str | None, program_id: str | None, class_id: str | None) -> bool:
    if scope is None:
        return True
    if not scope:
        return False
    for entry in scope:
        if entry.get("college") and entry["college"] != college_id:
            continue
        if entry.get("program") and entry["program"] != program_id:
            continue
        if entry.get("class") and entry["class"] != class_id:
            continue
        return True
    return False


def _ensure_class_in_scope(admin: dict, class_id: str) -> None:
    """超级管理员不受限；普通管理员操作的班级必须落在其数据范围内。"""
    scope = admin_student_scope(admin)
    if scope is None:
        return
    row = one(
        """
        SELECT classes.id AS class_id, classes.program_id AS program_id,
               programs.college_id AS college_id
        FROM campus_classes AS classes
        JOIN campus_programs AS programs ON programs.id = classes.program_id
        WHERE classes.id = ?
        """,
        (class_id,),
    )
    if not row or not scope_allows(scope, row.get("college_id"), row.get("program_id"), class_id):
        raise error(403, "该班级不在你管理的数据范围内。")


def _ensure_student_in_scope(admin: dict, user_id: str) -> None:
    """超级管理员不受限；普通管理员只能操作其数据范围内的学生。"""
    scope = admin_student_scope(admin)
    if scope is None:
        return
    row = one(
        """
        SELECT classes.id AS class_id, classes.program_id AS program_id,
               programs.college_id AS college_id
        FROM student_enrollments AS enrollments
        LEFT JOIN campus_classes AS classes ON classes.id = enrollments.class_id
        LEFT JOIN campus_programs AS programs ON programs.id = classes.program_id
        WHERE enrollments.user_id = ?
        """,
        (user_id,),
    )
    college_id = row.get("college_id") if row else None
    program_id = row.get("program_id") if row else None
    class_id = row.get("class_id") if row else None
    if not class_id or not scope_allows(scope, college_id, program_id, class_id):
        raise error(403, "该学生不在你管理的数据范围内。")


def require_super_admin(admin: dict = Depends(require_admin)) -> dict:
    if not is_super_admin(admin):
        raise error(403, "该操作仅限超级管理员执行。")
    return admin


def require_permission(*perms: str):
    """依赖注入：超级管理员放行；普通管理员需持有任意一个权限点。"""
    allowed = set(perms)

    def dependency(admin: dict = Depends(require_admin)) -> dict:
        if not is_super_admin(admin) and (not allowed or admin_permissions(admin).isdisjoint(allowed)):
            raise error(403, "当前管理员没有执行该操作的权限。")
        return admin

    return dependency


def require_catalog_permission(permission: str):
    def dependency(admin: dict = Depends(require_admin)) -> dict:
        has_new_permission = "manageCatalog" in admin_permissions(admin)
        has_legacy_permission = has_permission(
            db,
            str(admin.get("role") or ""),
            permission,
        )
        if not is_super_admin(admin) and not has_new_permission and not has_legacy_permission:
            raise error(403, "当前管理员没有岗位知识库管理权限。")
        return admin

    return dependency


def create_admin_session(response: Response, request: Request, admin_user_id: str) -> None:
    token = create_token()
    db.execute(
        """
        INSERT INTO admin_sessions (id, admin_user_id, token_hash, expires_at, user_agent, ip_address)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            str(uuid4()),
            admin_user_id,
            hash_token(token),
            utc_after(ADMIN_SESSION_MAX_AGE),
            request.headers.get("user-agent"),
            request_ip(request),
        ),
    )
    db.commit()
    response.set_cookie(
        ADMIN_SESSION_COOKIE,
        token,
        httponly=True,
        secure=os.environ.get("ADMIN_COOKIE_SECURE", "false").lower() == "true",
        samesite="lax",
        path="/",
        max_age=ADMIN_SESSION_MAX_AGE,
    )


def clear_admin_session(response: Response) -> None:
    response.delete_cookie(
        ADMIN_SESSION_COOKIE,
        httponly=True,
        secure=os.environ.get("ADMIN_COOKIE_SECURE", "false").lower() == "true",
        samesite="lax",
        path="/",
    )


def record_audit(
    request: Request,
    admin: dict | None,
    action: str,
    *,
    target_type: str = "system",
    target_id: str = "",
    summary: str = "",
    success: bool = True,
) -> None:
    db.execute(
        """
        INSERT INTO admin_audit_logs (
          id, admin_user_id, actor_email, action, target_type, target_id,
          summary, ip_address, user_agent, success
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            str(uuid4()),
            admin.get("id") if admin else None,
            admin.get("email") if admin else None,
            action,
            target_type,
            target_id or None,
            summary[:1000],
            request_ip(request),
            request.headers.get("user-agent"),
            1 if success else 0,
        ),
    )
    db.commit()


def login_attempt_key(request: Request, email: str) -> str:
    return f"{request_ip(request)}:{email}"


def is_login_limited(request: Request, email: str) -> bool:
    key = login_attempt_key(request, email)
    record = login_attempts.get(key)
    if not record:
        return False
    if float(record["reset_at"]) <= time.time():
        login_attempts.pop(key, None)
        return False
    return int(record["count"]) >= MAX_LOGIN_ATTEMPTS


def record_failed_login(request: Request, email: str) -> None:
    key = login_attempt_key(request, email)
    now = time.time()
    record = login_attempts.get(key)
    if not record or float(record["reset_at"]) <= now:
        login_attempts[key] = {"count": 1, "reset_at": now + LOGIN_WINDOW_SECONDS}
    else:
        record["count"] = int(record["count"]) + 1


def clear_failed_logins(request: Request, email: str) -> None:
    login_attempts.pop(login_attempt_key(request, email), None)


def bootstrap_master_admin() -> None:
    """启动时确保代码定义的主管理员存在：角色固定为 super_admin，权限点为全量。"""
    email = MASTER_ADMIN_EMAIL
    if not is_valid_email(email) or len(MASTER_ADMIN_PASSWORD) < 12:
        raise RuntimeError("主管理员配置无效：邮箱需有效且密码至少 12 位。")
    permissions = json.dumps(list(ALL_PERMISSIONS))
    existing = one("SELECT id FROM admin_users WHERE email = ?", (email,))
    if existing:
        db.execute(
            """
            UPDATE admin_users
            SET name = ?, role = 'super_admin', status = 'normal',
                password_hash = ?, permissions = ?, student_scope = NULL
            WHERE email = ?
            """,
            (MASTER_ADMIN_NAME, hash_password(MASTER_ADMIN_PASSWORD), permissions, email),
        )
        db.commit()
        return
    db.execute(
        """
        INSERT INTO admin_users (id, email, password_hash, name, role, status, permissions, student_scope)
        VALUES (?, ?, ?, ?, 'super_admin', 'normal', ?, NULL)
        """,
        (str(uuid4()), email, hash_password(MASTER_ADMIN_PASSWORD), MASTER_ADMIN_NAME, permissions),
    )
    db.commit()


def list_audit_logs(limit: int = 50) -> list[dict[str, Any]]:
    rows = all_rows(
        """
        SELECT id, actor_email, action, target_type, target_id, summary, success, created_at
        FROM admin_audit_logs
        ORDER BY created_at DESC
        LIMIT ?
        """,
        (max(1, min(limit, 100)),),
    )
    return [
        {
            "id": row["id"],
            "actor": row.get("actor_email") or "系统",
            "action": row["action"],
            "target": f"{row.get('target_type') or 'system'}:{row.get('target_id') or '-'}",
            "summary": row.get("summary") or "",
            "success": bool(row.get("success")),
            "time": str(row.get("created_at") or "-"),
        }
        for row in rows
    ]


def list_admin_accounts() -> list[dict[str, Any]]:
    users = all_rows(
        """
        SELECT id, email, name, role, status, permissions, student_scope, created_at, updated_at, last_login_at
        FROM admin_users ORDER BY created_at ASC
        """
    )
    return [sanitize_admin(user) | {"created_at": str(user.get("created_at") or "-")} for user in users]


ensure_admin_schema()
bootstrap_master_admin()


def count_value(sql: str, params: tuple = ()) -> int:
    row = one(sql, params) or {}
    return int(row.get("count") or 0)


def avg_value(sql: str, params: tuple = ()) -> float:
    row = one(sql, params) or {}
    return round(float(row.get("value") or 0), 1)


def status_label(value: str | None) -> str:
    return {"draft": "草稿", "running": "进行中", "completed": "已完成", "normal": "正常", "disabled": "已禁用"}.get(value or "", value or "-")


def report_review_status(score: int) -> str:
    if score < 70:
        return "需要复核"
    if score < 80:
        return "待抽检"
    return "已复核"


def recommendation_label(value: str | None) -> str:
    return {
        "strong_pass": "准备充分",
        "pass": "基本准备",
        "next_round": "建议进阶训练",
        "hold": "建议继续练习",
        "reject": "建议重点提升",
    }.get(value or "", value or "-")


def read_env_file(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.exists():
        return values
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


PLACEHOLDER_VALUES = {
    "your-api-key-here",
    "your-dashscope-api-key",
    "your-kimi-api-key-here",
    "your-deepseek-api-key-here",
    "placeholder",
}


def is_placeholder_secret(value: str | None) -> bool:
    cleaned = (value or "").strip().lower()
    if not cleaned:
        return False
    return (
        cleaned in PLACEHOLDER_VALUES
        or cleaned.startswith("your-")
        or "api-key-here" in cleaned
    )


def configured_secret(env_values: dict[str, str], *keys: str) -> str:
    for key in keys:
        value = (env_values.get(key) or os.environ.get(key, "")).strip()
        if value and not is_placeholder_secret(value):
            return value
    return ""


def write_env_values(path: Path, updates: dict[str, str]) -> None:
    existing_lines = path.read_text(encoding="utf-8").splitlines() if path.exists() else []
    seen: set[str] = set()
    next_lines: list[str] = []
    for raw_line in existing_lines:
        stripped = raw_line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            next_lines.append(raw_line)
            continue
        key, _value = raw_line.split("=", 1)
        key = key.strip()
        if key in updates:
            next_lines.append(f"{key}={updates[key]}")
            seen.add(key)
        else:
            next_lines.append(raw_line)
    for key, value in updates.items():
        if key not in seen:
            next_lines.append(f"{key}={value}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(next_lines).rstrip() + "\n", encoding="utf-8")


def mask_secret(value: str | None) -> str:
    if not value:
        return ""
    if len(value) <= 8:
        return "已配置"
    return f"{value[:4]}...{value[-4:]}"


def build_settings() -> dict[str, Any]:
    backend_env = read_env_file(USER_BACKEND_ENV_PATH)
    openai_key = configured_secret(backend_env, "OPENAI_REALTIME_API_KEY", "OPENAI_API_KEY")
    report_openai_key = configured_secret(backend_env, "OPENAI_API_KEY")
    report_qwen_key = configured_secret(backend_env, "QWEN_API_KEY")
    dashscope_key = configured_secret(backend_env, "DASHSCOPE_API_KEY")
    qwen_omni_key = configured_secret(
        backend_env,
        "QWEN_OMNI_REALTIME_API_KEY",
        "QWEN_OMNI_API_KEY",
        "DASHSCOPE_API_KEY",
    )
    return {
        "openaiRealtimeModel": backend_env.get("OPENAI_REALTIME_MODEL") or os.environ.get("OPENAI_REALTIME_MODEL", "gpt-realtime-2"),
        "openaiVoice": backend_env.get("OPENAI_REALTIME_VOICE") or os.environ.get("OPENAI_REALTIME_VOICE", "marin"),
        "qwenTtsModel": backend_env.get("DASHSCOPE_TTS_MODEL") or os.environ.get("DASHSCOPE_TTS_MODEL", "cosyvoice-v3-flash"),
        "qwenTtsVoice": backend_env.get("DASHSCOPE_TTS_VOICE") or os.environ.get("DASHSCOPE_TTS_VOICE", "longanyang"),
        "qwenTtsRegion": backend_env.get("DASHSCOPE_TTS_REGION") or os.environ.get("DASHSCOPE_TTS_REGION", "beijing"),
        "qwenTtsWorkspaceId": backend_env.get("DASHSCOPE_WORKSPACE_ID") or os.environ.get("DASHSCOPE_WORKSPACE_ID", ""),
        "qwenOmniModel": backend_env.get("QWEN_OMNI_REALTIME_MODEL") or os.environ.get("QWEN_OMNI_REALTIME_MODEL", "qwen3.5-omni-plus-realtime"),
        "qwenOmniVoice": backend_env.get("QWEN_OMNI_REALTIME_VOICE") or os.environ.get("QWEN_OMNI_REALTIME_VOICE", "Tina"),
        "qwenOmniRegion": backend_env.get("QWEN_OMNI_REALTIME_REGION") or os.environ.get("QWEN_OMNI_REALTIME_REGION", "beijing"),
        "qwenOmniWorkspaceId": backend_env.get("QWEN_OMNI_REALTIME_WORKSPACE_ID") or backend_env.get("DASHSCOPE_WORKSPACE_ID") or os.environ.get("QWEN_OMNI_REALTIME_WORKSPACE_ID") or os.environ.get("DASHSCOPE_WORKSPACE_ID", ""),
        "qwenOmniEndpoint": backend_env.get("QWEN_OMNI_REALTIME_WEBRTC_ENDPOINT") or os.environ.get("QWEN_OMNI_REALTIME_WEBRTC_ENDPOINT", ""),
        "reviewRule": "低于 70 分自动进入复核",
        "openaiApiKeyConfigured": bool(openai_key and openai_key != "your-api-key-here"),
        "openaiApiKeyMasked": mask_secret(openai_key),
        "reportOpenaiApiKeyConfigured": bool(report_openai_key),
        "reportOpenaiApiKeyMasked": mask_secret(report_openai_key),
        "reportOpenaiModel": backend_env.get("OPENAI_FOLLOWUP_MODEL") or os.environ.get("OPENAI_FOLLOWUP_MODEL", "gpt-4o-mini"),
        "reportQwenApiKeyConfigured": bool(report_qwen_key),
        "reportQwenApiKeyMasked": mask_secret(report_qwen_key),
        "reportQwenModel": backend_env.get("QWEN_MODEL") or os.environ.get("QWEN_MODEL", "qwen-plus"),
        "reportProviderOrder": backend_env.get("REPORT_PROVIDER_ORDER") or os.environ.get("REPORT_PROVIDER_ORDER", "openai,qwen"),
        "reportTimeout": int(backend_env.get("AI_REPORT_TIMEOUT") or os.environ.get("AI_REPORT_TIMEOUT", "60")),
        "reportRetries": int(backend_env.get("AI_REPORT_HTTP_RETRIES") or os.environ.get("AI_REPORT_HTTP_RETRIES", "1")),
        "dashscopeApiKeyConfigured": bool(dashscope_key),
        "dashscopeApiKeyMasked": mask_secret(dashscope_key),
        "qwenOmniApiKeyConfigured": bool(qwen_omni_key),
        "qwenOmniApiKeyMasked": mask_secret(qwen_omni_key),
        "envPath": str(USER_BACKEND_ENV_PATH),
    }


STUDENT_IMPORT_HEADERS = {
    "学院": "college",
    "学号": "student_no",
    "姓名": "name",
    "性别": "gender",
    "班级": "class_name",
    "辅导员": "counselor",
    "账号状态": "source_account_status",
    "学生状态": "student_status",
    "注册时间": "source_registered_at",
    "修改时间": "source_updated_at",
}
STUDENT_IMPORT_MAX_BYTES = 10 * 1024 * 1024
STUDENT_IMPORT_MAX_ROWS = 10000


def normalize_student_no(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value or "").strip().upper()


def excel_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def student_login_status(account_status: str, student_status: str) -> str:
    disabled_markers = ("禁用", "停用", "冻结", "注销", "退学", "无效", "disabled", "inactive")
    combined = f"{account_status} {student_status}".lower()
    return "disabled" if any(marker in combined for marker in disabled_markers) else "normal"


def student_account_email(student_no: str) -> str:
    return f"{student_no.lower()}@student.local"


def student_admission_year(student_no: object) -> str:
    matched = re.match(r"^(20\d{2})", str(student_no or "").strip())
    return matched.group(1) if matched else ""


def import_student_accounts_from_workbook(content: bytes) -> dict[str, Any]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise error(400, "无法读取 Excel，请确认文件为有效的 .xlsx 格式。") from exc

    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise error(400, "Excel 中没有可导入的数据。") from exc

    headers = [excel_cell_text(value) for value in raw_headers]
    column_indexes = {
        field: index
        for index, header in enumerate(headers)
        if (field := STUDENT_IMPORT_HEADERS.get(header))
    }
    missing = [label for label, field in (("学号", "student_no"), ("姓名", "name")) if field not in column_indexes]
    if missing:
        raise error(400, f"Excel 缺少必填列：{'、'.join(missing)}。")

    created = 0
    updated = 0
    migrated = 0
    skipped = 0
    errors: list[str] = []
    seen_student_numbers: set[str] = set()

    db.begin()
    try:
        for row_number, raw_row in enumerate(rows, start=2):
            if row_number > STUDENT_IMPORT_MAX_ROWS + 1:
                errors.append(f"最多支持导入 {STUDENT_IMPORT_MAX_ROWS} 名学生，其余行未处理。")
                break
            values = {
                field: excel_cell_text(raw_row[index] if index < len(raw_row) else None)
                for field, index in column_indexes.items()
            }
            if not any(values.values()):
                continue
            student_no = normalize_student_no(values.get("student_no"))
            name = values.get("name", "").strip()
            if not re.match(r"^[0-9A-Z_-]{3,40}$", student_no):
                skipped += 1
                errors.append(f"第 {row_number} 行学号格式不正确。")
                continue
            if not name or len(name) > 120:
                skipped += 1
                errors.append(f"第 {row_number} 行姓名为空或过长。")
                continue
            if student_no in seen_student_numbers:
                skipped += 1
                errors.append(f"第 {row_number} 行学号 {student_no} 在文件中重复。")
                continue
            seen_student_numbers.add(student_no)

            account_status = values.get("source_account_status", "")[:80]
            student_status = values.get("student_status", "")[:80]
            login_status = student_login_status(account_status, student_status)
            existing = one("SELECT id FROM users WHERE student_no = ?", (student_no,))
            legacy_registration = one(
                "SELECT id, user_id FROM student_registrations WHERE student_no = ?",
                (student_no,),
            )
            common_values = (
                name,
                values.get("college", "")[:160] or None,
                values.get("gender", "")[:20] or None,
                values.get("class_name", "")[:160] or None,
                values.get("counselor", "")[:120] or None,
                student_status or None,
                account_status or None,
                values.get("source_registered_at") or None,
                values.get("source_updated_at") or None,
                login_status,
            )
            if existing:
                cursor = db.execute(
                    """
                    UPDATE users
                    SET name = ?, college = ?, gender = ?, class_name = ?, counselor = ?,
                        student_status = ?, source_account_status = ?, source_registered_at = ?,
                        source_updated_at = ?, status = ?, updated_at = CURRENT_TIMESTAMP
                    WHERE id = ?
                    """,
                    common_values + (existing["id"],),
                )
                cursor.close()
                if one("SELECT user_id FROM profiles WHERE user_id = ?", (existing["id"],)):
                    cursor = db.execute("UPDATE profiles SET nickname = ?, updated_at = CURRENT_TIMESTAMP WHERE user_id = ?", (name, existing["id"]))
                    cursor.close()
                else:
                    cursor = db.execute(
                        "INSERT INTO profiles (id, user_id, nickname) VALUES (?, ?, ?)",
                        (str(uuid4()), existing["id"], name),
                    )
                    cursor.close()
                if login_status != "normal":
                    cursor = db.execute("DELETE FROM sessions WHERE user_id = ?", (existing["id"],))
                    cursor.close()
                updated += 1
                continue

            legacy_user = None
            if legacy_registration and legacy_registration.get("user_id"):
                legacy_user = one(
                    "SELECT id, student_no FROM users WHERE id = ?",
                    (legacy_registration["user_id"],),
                )
            if not legacy_user:
                legacy_user = one(
                    "SELECT id, student_no FROM users WHERE email = ?",
                    (student_account_email(student_no),),
                )
            if legacy_user:
                if legacy_user.get("student_no") not in (None, "", student_no):
                    skipped += 1
                    errors.append(f"第 {row_number} 行学号 {student_no} 对应的历史账号已绑定其他学号，需人工核对。")
                    continue
                if legacy_user.get("student_no") in (None, "", student_no):
                    temporary_password = generate_temporary_password()
                    cursor = db.execute(
                        """
                        UPDATE users
                        SET student_no = ?, password_hash = ?, must_change_password = 1,
                            temp_password_encrypted = ?, temp_password_created_at = CURRENT_TIMESTAMP,
                            name = ?, college = ?, gender = ?, class_name = ?, counselor = ?,
                            student_status = ?, source_account_status = ?, source_registered_at = ?,
                            source_updated_at = ?, status = ?, updated_at = CURRENT_TIMESTAMP
                        WHERE id = ?
                        """,
                        (student_no, hash_password(temporary_password),
                         encrypt_temporary_password(temporary_password))
                        + common_values + (legacy_user["id"],),
                    )
                    cursor.close()
                    cursor = db.execute("DELETE FROM sessions WHERE user_id = ?", (legacy_user["id"],))
                    cursor.close()
                    if legacy_registration and legacy_registration.get("user_id") != legacy_user["id"]:
                        cursor = db.execute(
                            "UPDATE student_registrations SET user_id = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                            (legacy_user["id"], legacy_registration["id"]),
                        )
                        cursor.close()
                    if one("SELECT user_id FROM profiles WHERE user_id = ?", (legacy_user["id"],)):
                        cursor = db.execute(
                            "UPDATE profiles SET nickname = ?, updated_at = CURRENT_TIMESTAMP WHERE user_id = ?",
                            (name, legacy_user["id"]),
                        )
                    else:
                        cursor = db.execute(
                            "INSERT INTO profiles (id, user_id, nickname) VALUES (?, ?, ?)",
                            (str(uuid4()), legacy_user["id"], name),
                        )
                    cursor.close()
                    migrated += 1
                    continue

            temporary_password = generate_temporary_password()
            user_id = str(uuid4())
            cursor = db.execute(
                """
                INSERT INTO users (
                  id, email, student_no, password_hash, name, college, gender, class_name,
                  counselor, student_status, source_account_status, must_change_password,
                  temp_password_encrypted, temp_password_created_at, source_registered_at,
                  source_updated_at, status
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, CURRENT_TIMESTAMP, ?, ?, ?)
                """,
                (
                    user_id,
                    student_account_email(student_no),
                    student_no,
                    hash_password(temporary_password),
                ) + common_values[:7] + (encrypt_temporary_password(temporary_password),) + common_values[7:9] + (login_status,),
            )
            cursor.close()
            if legacy_registration:
                cursor = db.execute(
                    "UPDATE student_registrations SET user_id = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                    (user_id, legacy_registration["id"]),
                )
                cursor.close()
            cursor = db.execute(
                "INSERT INTO profiles (id, user_id, nickname) VALUES (?, ?, ?)",
                (str(uuid4()), user_id, name),
            )
            cursor.close()
            created += 1
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        workbook.close()

    return {
        "created": created,
        "updated": updated,
        "migrated": migrated,
        "skipped": skipped,
        "processed": created + updated + migrated,
        "errors": errors[:30],
    }


def build_student_password_workbook(rows: list[dict[str, Any]]) -> io.BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "学生临时密码"
    headers = ["学院", "学号", "学年", "姓名", "性别", "班级", "辅导员", "临时密码", "账号状态", "学生状态", "激活状态", "生成时间"]
    worksheet.append(headers)
    header_fill = PatternFill("solid", fgColor="17324D")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        worksheet.append([
            row.get("college") or "",
            row.get("student_no") or "",
            student_admission_year(row.get("student_no")),
            row.get("name") or "",
            row.get("gender") or "",
            row.get("class_name") or "",
            row.get("counselor") or "",
            decrypt_temporary_password(row["temp_password_encrypted"]),
            row.get("source_account_status") or row.get("status") or "",
            row.get("student_status") or "",
            "待首次改密",
            str(row.get("temp_password_created_at") or ""),
        ])

    widths = [18, 18, 10, 14, 9, 18, 14, 18, 14, 14, 14, 20]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "1:1"
    worksheet.oddFooter.center.text = "仅限辅导员发放临时密码使用，请妥善保管"

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def list_candidates(scope: list[dict[str, str]] | None = None) -> list[dict[str, Any]]:
    rows = all_rows(
        """
        SELECT users.id, users.email, users.student_no, users.name, users.college, users.class_name,
               users.counselor, users.student_status, users.must_change_password,
               users.temp_password_encrypted, users.status, users.created_at, users.last_login_at,
               profiles.target_role,
               enrollments.class_id AS class_id, classes.program_id AS program_id,
               programs.college_id AS college_id,
               COALESCE(interview_counts.interviews, 0) AS interviews,
               COALESCE(report_scores.average_score, 0) AS average_score
        FROM users
        LEFT JOIN profiles ON profiles.user_id = users.id
        LEFT JOIN student_enrollments AS enrollments ON enrollments.user_id = users.id
        LEFT JOIN campus_classes AS classes ON classes.id = enrollments.class_id
        LEFT JOIN campus_programs AS programs ON programs.id = classes.program_id
        LEFT JOIN (
          SELECT user_id, COUNT(*) AS interviews
          FROM interview_sessions
          GROUP BY user_id
        ) AS interview_counts ON interview_counts.user_id = users.id
        LEFT JOIN (
          SELECT user_id, AVG(total_score) AS average_score
          FROM interview_reports
          GROUP BY user_id
        ) AS report_scores ON report_scores.user_id = users.id
        ORDER BY users.updated_at DESC, users.created_at DESC
        LIMIT 5000
        """
    )
    result = []

    for row in rows:
        if scope is not None and not scope_allows(
            scope,
            row.get("college_id"),
            row.get("program_id"),
            row.get("class_id"),
        ):
            continue

        result.append(
            {
                "id": row["id"],
                "name": row["name"],
                "email": row["email"],
                "studentNo": row.get("student_no") or "-",
                "admissionYear": student_admission_year(row.get("student_no") or ""),
                "college": row.get("college") or "-",
                "className": row.get("class_name") or "-",
                "counselor": row.get("counselor") or "-",
                "role": row.get("target_role") or "未填写",
                "status": status_label(row.get("status")),
                "activationStatus": (
                    "准备改密"
                    if row.get("must_change_password")
                    else ("已激活" if row.get("student_no") else "待绑定学号")
                ),
                "canViewTemporaryPassword": bool(
                    row.get("must_change_password")
                    and row.get("temp_password_encrypted")
                ),
                "interviews": int(row.get("interviews") or 0),
                "averageScore": round(float(row.get("average_score") or 0), 1),
                "lastLogin": str(row.get("last_login_at") or "-"),
            }
        )

    return result


def get_candidate_detail(candidate_id: str) -> dict[str, Any]:
    row = one(
        """
        SELECT users.id, users.email, users.student_no, users.name, users.college, users.gender,
               users.class_name, users.counselor, users.student_status, users.source_account_status,
               users.must_change_password,
               CASE WHEN users.temp_password_encrypted IS NULL THEN 0 ELSE 1 END AS can_view_temp_password,
               users.temp_password_created_at,
               users.activated_at, users.status, users.created_at, users.updated_at,
               users.last_login_at, profiles.nickname, profiles.target_role, profiles.experience_level,
               profiles.company_type, profiles.target_city, profiles.expected_salary,
               profiles.years_of_experience, profiles.education_level, profiles.skills,
               profiles.project_keywords, profiles.resume_text, profiles.project_experience,
               profiles.portfolio_links
        FROM users
        LEFT JOIN profiles ON profiles.user_id = users.id
        WHERE users.id = ?
        """,
        (candidate_id,),
    )
    if not row:
        raise error(404, "用户不存在。")
    return row


def list_interviews() -> list[dict[str, Any]]:
    rows = all_rows(
        """
        SELECT interviews.id, interviews.target_role, interviews.interview_type, interviews.status,
               interviews.updated_at, users.name AS candidate,
               COALESCE(agent_counts.agent_count, 0) AS agent_count,
               COALESCE(message_counts.message_count, 0) AS message_count
        FROM interview_sessions AS interviews
        JOIN users ON users.id = interviews.user_id
        LEFT JOIN (
          SELECT interview_id, COUNT(*) AS agent_count
          FROM interview_agents
          GROUP BY interview_id
        ) AS agent_counts ON agent_counts.interview_id = interviews.id
        LEFT JOIN (
          SELECT interview_id, COUNT(*) AS message_count
          FROM interview_messages
          GROUP BY interview_id
        ) AS message_counts ON message_counts.interview_id = interviews.id
        ORDER BY interviews.updated_at DESC, interviews.created_at DESC
        LIMIT 50
        """
    )
    return [
        {
            "id": row["id"],
            "candidate": row["candidate"],
            "role": row.get("target_role") or "-",
            "type": row.get("interview_type") or "综合模拟",
            "status": status_label(row.get("status")),
            "agents": f"{int(row.get('agent_count') or 0)}/3",
            "messages": int(row.get("message_count") or 0),
            "updatedAt": str(row.get("updated_at") or "-"),
        }
        for row in rows
    ]


def list_connection_logs(
    *,
    level: str = "",
    event_type: str = "",
    query: str = "",
    limit: int = 300,
) -> list[dict[str, Any]]:
    conditions: list[str] = []
    params: list[Any] = []
    if level:
        conditions.append("events.level = ?")
        params.append(level)
    if event_type:
        conditions.append("events.event_type = ?")
        params.append(event_type)
    if query:
        conditions.append(
            """
            (
              users.name LIKE ? OR interviews.target_role LIKE ?
              OR events.session_id LIKE ? OR events.interview_id LIKE ?
              OR events.event_type LIKE ? OR events.message LIKE ?
            )
            """
        )
        keyword = f"%{query}%"
        params.extend([keyword] * 6)
    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    safe_limit = max(1, min(limit, 500))
    params.append(safe_limit)
    rows = all_rows(
        f"""
        SELECT events.id, events.interview_id, events.session_id, events.provider,
               events.event_type, events.level, events.connection_state,
               events.ice_connection_state, events.ice_gathering_state,
               events.signaling_state, events.data_channel_state, events.message,
               events.metadata_json, events.client_created_at, events.created_at,
               users.name AS candidate, interviews.target_role
        FROM webrtc_diagnostic_events AS events
        JOIN interview_sessions AS interviews ON interviews.id = events.interview_id
        JOIN users ON users.id = events.user_id
        {where_clause}
        ORDER BY events.created_at DESC
        LIMIT ?
        """,
        tuple(params),
    )
    result: list[dict[str, Any]] = []
    for row in rows:
        try:
            metadata = json.loads(row.get("metadata_json") or "{}")
        except (TypeError, ValueError):
            metadata = {}
        result.append(
            {
                "id": row["id"],
                "interviewId": row["interview_id"],
                "sessionId": row.get("session_id") or "-",
                "candidate": row.get("candidate") or "未命名学生",
                "targetRole": row.get("target_role") or "未填写",
                "provider": row.get("provider") or "-",
                "eventType": row.get("event_type") or "-",
                "level": row.get("level") or "info",
                "connectionState": row.get("connection_state") or "-",
                "iceConnectionState": row.get("ice_connection_state") or "-",
                "iceGatheringState": row.get("ice_gathering_state") or "-",
                "signalingState": row.get("signaling_state") or "-",
                "dataChannelState": row.get("data_channel_state") or "-",
                "message": row.get("message") or "",
                "metadata": metadata,
                "clientCreatedAt": str(row.get("client_created_at") or "-"),
                "createdAt": str(row.get("created_at") or "-"),
            }
        )
    return result


def connection_log_summary() -> dict[str, Any]:
    recent_clause = (
        "created_at >= DATE_SUB(CURRENT_TIMESTAMP, INTERVAL 1 DAY)"
        if DB_ENGINE == "mysql"
        else "created_at >= datetime('now', '-1 day')"
    )
    row = one(
        f"""
        SELECT COUNT(*) AS total,
               SUM(CASE WHEN level = 'error' THEN 1 ELSE 0 END) AS errors,
               SUM(CASE WHEN level = 'warning' THEN 1 ELSE 0 END) AS warnings,
               COUNT(DISTINCT session_id) AS sessions,
               MAX(created_at) AS latest_at
        FROM webrtc_diagnostic_events
        WHERE {recent_clause}
        """
    ) or {}
    return {
        "total": int(row.get("total") or 0),
        "errors": int(row.get("errors") or 0),
        "warnings": int(row.get("warnings") or 0),
        "sessions": int(row.get("sessions") or 0),
        "latestAt": str(row.get("latest_at") or "-"),
    }


def get_interview_detail(interview_id: str) -> dict[str, Any]:
    interview = one(
        """
        SELECT interviews.*, users.name AS candidate, users.email AS candidate_email
        FROM interview_sessions AS interviews
        JOIN users ON users.id = interviews.user_id
        WHERE interviews.id = ?
        """,
        (interview_id,),
    )
    if not interview:
        raise error(404, "面试不存在。")
    agents = all_rows(
        """
        SELECT id, agent_name, agent_type, agent_role, strategy, order_index, status, created_at, updated_at
        FROM interview_agents
        WHERE interview_id = ?
        ORDER BY order_index ASC, created_at ASC
        """,
        (interview_id,),
    )
    messages = all_rows(
        """
        SELECT messages.id, messages.sender_type, messages.message_type, messages.content,
               messages.transcript_text, messages.order_index, messages.created_at,
               agents.agent_name, agents.agent_type
        FROM interview_messages AS messages
        LEFT JOIN interview_agents AS agents ON agents.id = messages.agent_id
        WHERE messages.interview_id = ?
        ORDER BY messages.order_index ASC, messages.created_at ASC
        LIMIT 80
        """,
        (interview_id,),
    )
    return {"interview": interview, "agents": agents, "messages": messages}


def list_reports() -> list[dict[str, Any]]:
    rows = all_rows(
        """
        SELECT reports.id, reports.total_score, reports.grade, reports.pass_recommendation,
               reports.review_status, reports.provider, reports.model, reports.fallback,
               reports.created_at, users.name AS candidate, interviews.target_role
        FROM interview_reports AS reports
        JOIN users ON users.id = reports.user_id
        JOIN interview_sessions AS interviews ON interviews.id = reports.interview_id
        ORDER BY reports.updated_at DESC, reports.created_at DESC
        LIMIT 50
        """
    )
    return [
        {
            "id": row["id"],
            "candidate": row["candidate"],
            "role": row.get("target_role") or "-",
            "score": int(row.get("total_score") or 0),
            "grade": row.get("grade") or "-",
            "recommendation": recommendation_label(row.get("pass_recommendation")),
            "reviewStatus": {"approved": "已复核", "rejected": "复核未通过", "pending": "待复核"}.get(row.get("review_status"), "待复核"),
            "provider": row.get("provider") or "local",
            "model": row.get("model") or "rules-v1",
            "fallback": bool(row.get("fallback")),
            "createdAt": str(row.get("created_at") or "-"),
        }
        for row in rows
    ]


def get_report_detail(report_id: str) -> dict[str, Any]:
    report = one(
        """
        SELECT reports.*, users.name AS candidate, users.email AS candidate_email,
               interviews.target_role, interviews.interview_type, interviews.difficulty
        FROM interview_reports AS reports
        JOIN users ON users.id = reports.user_id
        JOIN interview_sessions AS interviews ON interviews.id = reports.interview_id
        WHERE reports.id = ?
        """,
        (report_id,),
    )
    if not report:
        raise error(404, "报告不存在。")
    return report


def list_agent_templates() -> list[dict[str, Any]]:
    rows = all_rows(
        """
        SELECT agent_name, agent_type, agent_role, COUNT(*) AS usage_count
        FROM interview_agents
        GROUP BY agent_name, agent_type, agent_role
        ORDER BY usage_count DESC, agent_name ASC
        LIMIT 20
        """
    )
    return [
        {
            "name": row.get("agent_name") or "-",
            "type": row.get("agent_type") or "-",
            "status": "启用",
            "focus": row.get("agent_role") or "-",
            "usageCount": int(row.get("usage_count") or 0),
        }
        for row in rows
    ]


def get_agent_detail(agent_name: str) -> dict[str, Any]:
    row = one(
        """
        SELECT agent_name, agent_type, agent_role, strategy, COUNT(*) AS usage_count
        FROM interview_agents
        WHERE agent_name = ?
        GROUP BY agent_name, agent_type, agent_role, strategy
        ORDER BY usage_count DESC
        LIMIT 1
        """,
        (agent_name,),
    )
    if not row:
        raise error(404, "Agent 不存在。")
    recent = all_rows(
        """
        SELECT agents.interview_id, agents.status, agents.updated_at,
               interviews.target_role, users.name AS candidate
        FROM interview_agents AS agents
        JOIN interview_sessions AS interviews ON interviews.id = agents.interview_id
        JOIN users ON users.id = interviews.user_id
        WHERE agents.agent_name = ?
        ORDER BY agents.updated_at DESC
        LIMIT 20
        """,
        (agent_name,),
    )
    return {"agent": row, "recentUsage": recent}


def _required_text(body: dict[str, Any], key: str, label: str, max_length: int = 160) -> str:
    value = str(body.get(key) or "").strip()
    if not value:
        raise error(400, f"{label}不能为空。")
    return value[:max_length]


def _campus_catalog_options() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    majors = all_rows(
        """
        SELECT majors.id, majors.code, majors.name, colleges.name AS college_name
        FROM catalog_majors AS majors
        JOIN catalog_colleges AS colleges ON colleges.id = majors.college_id
        JOIN catalog_versions AS versions ON versions.id = majors.version_id
        WHERE versions.status = 'published' AND majors.enabled = 1 AND colleges.enabled = 1
        ORDER BY colleges.sort_order, colleges.name, majors.sort_order, majors.name
        """
    )
    jobs = all_rows(
        """
        SELECT jobs.id, jobs.code, jobs.name
        FROM catalog_job_roles AS jobs
        JOIN catalog_versions AS versions ON versions.id = jobs.version_id
        WHERE versions.status = 'published' AND jobs.enabled = 1
        ORDER BY jobs.sort_order, jobs.name
        """
    )
    return majors, jobs


def _student_growth_status(row: dict[str, Any]) -> str:
    interviews = int(row.get("interviews") or 0)
    readiness = float(row.get("readiness") or 0)
    if bool(row.get("focus_flag")):
        return "重点关注"
    if not row.get("enrollment_id"):
        return "未归班"
    if interviews == 0:
        return "尚未训练"
    if interviews >= 2 and readiness < 60:
        return "需要关注"
    if interviews >= 2 and readiness >= 80:
        return "表现稳定"
    return "训练中"


def list_campus_students() -> list[dict[str, Any]]:
    rows = all_rows(
        """
        SELECT users.id, users.email, users.name, users.status, users.last_login_at,
               profiles.target_role,
               enrollments.id AS enrollment_id, enrollments.student_no,
               enrollments.status AS enrollment_status, enrollments.focus_flag, enrollments.note,
               classes.id AS class_id, classes.name AS class_name, classes.graduation_year,
               programs.id AS program_id, programs.name AS program_name, programs.direction,
               colleges.id AS college_id, colleges.name AS college_name,
               COALESCE(interview_counts.interviews, 0) AS interviews,
               COALESCE(report_scores.readiness, 0) AS readiness
        FROM users
        LEFT JOIN profiles ON profiles.user_id = users.id
        LEFT JOIN student_enrollments AS enrollments ON enrollments.user_id = users.id
        LEFT JOIN campus_classes AS classes ON classes.id = enrollments.class_id
        LEFT JOIN campus_programs AS programs ON programs.id = classes.program_id
        LEFT JOIN campus_colleges AS colleges ON colleges.id = programs.college_id
        LEFT JOIN (
          SELECT user_id, COUNT(*) AS interviews
          FROM interview_sessions
          GROUP BY user_id
        ) AS interview_counts ON interview_counts.user_id = users.id
        LEFT JOIN (
          SELECT user_id, AVG(total_score) AS readiness
          FROM interview_reports
          GROUP BY user_id
        ) AS report_scores ON report_scores.user_id = users.id
        ORDER BY enrollments.focus_flag DESC, users.updated_at DESC, users.created_at DESC
        LIMIT 500
        """
    )
    result = []
    for row in rows:
        item = {
            "id": row["id"],
            "name": row.get("name") or "未命名学生",
            "email": row.get("email") or "-",
            "studentNo": row.get("student_no") or "-",
            "collegeId": row.get("college_id"),
            "college": row.get("college_name") or "未归属",
            "programId": row.get("program_id"),
            "program": row.get("program_name") or "未归属",
            "direction": row.get("direction") or "",
            "classId": row.get("class_id"),
            "className": row.get("class_name") or "未归班",
            "graduationYear": row.get("graduation_year"),
            "targetRole": row.get("target_role") or "尚未选择",
            "readiness": round(float(row.get("readiness") or 0), 1),
            "interviews": int(row.get("interviews") or 0),
            "focus": bool(row.get("focus_flag")),
            "note": row.get("note") or "",
            "accountStatus": status_label(row.get("status")),
            "lastLogin": str(row.get("last_login_at") or "-"),
        }
        item["growthStatus"] = _student_growth_status(row)
        result.append(item)
    return result


def campus_overview_data(scope: list[dict[str, str]] | None = None) -> dict[str, Any]:
    colleges = all_rows("SELECT id, code, name, status, created_at FROM campus_colleges ORDER BY name")
    programs = all_rows(
        """
        SELECT programs.id, programs.college_id, programs.standard_major_code, programs.name,
               programs.direction, programs.coordinator, programs.status, colleges.name AS college_name
        FROM campus_programs AS programs
        JOIN campus_colleges AS colleges ON colleges.id = programs.college_id
        ORDER BY colleges.name, programs.name
        """
    )
    classes = all_rows(
        """
        SELECT classes.id, classes.program_id, classes.name, classes.graduation_year,
               classes.advisor, classes.invite_code, classes.status,
               programs.name AS program_name, programs.college_id
        FROM campus_classes AS classes
        JOIN campus_programs AS programs ON programs.id = classes.program_id
        ORDER BY classes.graduation_year DESC, classes.name
        """
    )
    mappings = all_rows(
        """
        SELECT mappings.program_id, mappings.job_role_id, mappings.priority, jobs.name AS job_name
        FROM program_job_roles AS mappings
        LEFT JOIN catalog_job_roles AS jobs ON jobs.id = mappings.job_role_id
        ORDER BY mappings.priority DESC, jobs.name
        """
    )
    majors, jobs = _campus_catalog_options()
    try:
        students = list_campus_students()
    except Exception:
        students = []
    if scope is not None:
        students = [
            item
            for item in students
            if scope_allows(scope, item.get("collegeId"), item.get("programId"), item.get("classId"))
        ]

    program_counts: dict[str, int] = {}
    class_counts: dict[str, int] = {}
    for student in students:
        if student.get("programId"):
            program_counts[student["programId"]] = program_counts.get(student["programId"], 0) + 1
        if student.get("classId"):
            class_counts[student["classId"]] = class_counts.get(student["classId"], 0) + 1
    for college in colleges:
        college["studentCount"] = len([item for item in students if item.get("collegeId") == college["id"]])
    for program in programs:
        program["studentCount"] = program_counts.get(program["id"], 0)
        program["jobs"] = [item for item in mappings if item["program_id"] == program["id"]]
    for class_item in classes:
        class_item["studentCount"] = class_counts.get(class_item["id"], 0)
    return {
        "colleges": colleges,
        "programs": programs,
        "classes": classes,
        "students": students,
        "standardMajors": majors,
        "jobRoles": jobs,
        "summary": {
            "colleges": len(colleges),
            "programs": len(programs),
            "classes": len(classes),
            "students": len(students),
            "unassigned": len([item for item in students if not item.get("classId")]),
            "focus": len([item for item in students if item.get("focus")]),
        },
    }


def upsert_student_enrollment(
    user_id: str,
    *,
    class_id: str | None,
    student_no: str = "",
    status: str = "active",
    focus_flag: bool = False,
    note: str = "",
) -> dict[str, Any]:
    if not one("SELECT id FROM users WHERE id = ?", (user_id,)):
        raise error(404, "学生账号不存在。")
    if class_id and not one("SELECT id FROM campus_classes WHERE id = ?", (class_id,)):
        raise error(404, "班级不存在。")
    current = one("SELECT id FROM student_enrollments WHERE user_id = ?", (user_id,))
    if current:
        db.execute(
            """
            UPDATE student_enrollments
            SET class_id = ?, student_no = ?, status = ?, focus_flag = ?, note = ?, updated_at = CURRENT_TIMESTAMP
            WHERE user_id = ?
            """,
            (class_id, student_no[:80] or None, status, 1 if focus_flag else 0, note[:1000] or None, user_id),
        )
    else:
        db.execute(
            """
            INSERT INTO student_enrollments (id, user_id, class_id, student_no, status, focus_flag, note)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (str(uuid4()), user_id, class_id, student_no[:80] or None, status, 1 if focus_flag else 0, note[:1000] or None),
        )
    db.commit()
    return one("SELECT * FROM student_enrollments WHERE user_id = ?", (user_id,)) or {}


def build_metrics(reports: list[dict[str, Any]]) -> list[dict[str, str]]:
    users = count_value("SELECT COUNT(*) AS count FROM users")
    interviews = count_value("SELECT COUNT(*) AS count FROM interview_sessions")
    completed = count_value("SELECT COUNT(*) AS count FROM interview_sessions WHERE status = ?", ("completed",))
    report_count = count_value("SELECT COUNT(*) AS count FROM interview_reports")
    average_score = avg_value("SELECT AVG(total_score) AS value FROM interview_reports")
    review_count = len([report for report in reports if report["reviewStatus"] != "已复核"])
    completion_rate = round(completed / interviews * 100) if interviews else 0
    return [
        {"label": "注册用户", "value": f"{users:,}", "note": "来自 users", "tone": "blue"},
        {"label": "面试任务", "value": f"{interviews:,}", "note": f"完成率 {completion_rate}%", "tone": "green"},
        {"label": "报告生成", "value": f"{report_count:,}", "note": f"平均分 {average_score}", "tone": "amber"},
        {"label": "待复核", "value": f"{review_count:,}", "note": "按分数规则计算", "tone": "red"},
    ]


@app.get("/api/admin/campus/overview")
def campus_overview(admin: dict = Depends(require_permission("manageStudents"))):
    return campus_overview_data(scope=admin_student_scope(admin))


@app.post("/api/admin/campus/colleges", status_code=201)
async def create_campus_college(request: Request, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    body = await request.json()
    code = _required_text(body, "code", "学院编码", 64).upper()
    name = _required_text(body, "name", "学院名称")
    if one("SELECT id FROM campus_colleges WHERE code = ? OR name = ?", (code, name)):
        raise error(409, "学院编码或名称已存在。")
    college_id = str(uuid4())
    db.execute(
        "INSERT INTO campus_colleges (id, code, name, status) VALUES (?, ?, ?, 'active')",
        (college_id, code, name),
    )
    db.commit()
    record_audit(request, admin, "campus.college.create", target_type="campus_college", target_id=college_id, summary=f"创建学院：{name}")
    return {"college": one("SELECT * FROM campus_colleges WHERE id = ?", (college_id,))}


@app.post("/api/admin/campus/programs", status_code=201)
async def create_campus_program(request: Request, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    body = await request.json()
    college_id = _required_text(body, "collegeId", "所属学院", 36)
    if not one("SELECT id FROM campus_colleges WHERE id = ?", (college_id,)):
        raise error(404, "所属学院不存在。")
    name = _required_text(body, "name", "专业名称")
    direction = str(body.get("direction") or "").strip()[:160]
    coordinator = str(body.get("coordinator") or "").strip()[:120]
    standard_major_code = str(body.get("standardMajorCode") or "").strip()[:64]
    if one(
        "SELECT id FROM campus_programs WHERE college_id = ? AND name = ? AND COALESCE(direction, '') = ?",
        (college_id, name, direction),
    ):
        raise error(409, "该学院下已经存在相同专业和培养方向。")
    program_id = str(uuid4())
    db.execute(
        """
        INSERT INTO campus_programs (id, college_id, standard_major_code, name, direction, coordinator, status)
        VALUES (?, ?, ?, ?, ?, ?, 'active')
        """,
        (program_id, college_id, standard_major_code or None, name, direction or None, coordinator or None),
    )
    db.commit()
    record_audit(request, admin, "campus.program.create", target_type="campus_program", target_id=program_id, summary=f"创建专业：{name}")
    return {"program": one("SELECT * FROM campus_programs WHERE id = ?", (program_id,))}


@app.post("/api/admin/campus/classes", status_code=201)
async def create_campus_class(request: Request, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    body = await request.json()
    program_id = _required_text(body, "programId", "所属专业", 36)
    if not one("SELECT id FROM campus_programs WHERE id = ?", (program_id,)):
        raise error(404, "所属专业不存在。")
    name = _required_text(body, "name", "班级名称")
    advisor = str(body.get("advisor") or "").strip()[:120]
    raw_year = body.get("graduationYear")
    try:
        graduation_year = int(raw_year) if raw_year not in {None, ""} else None
    except (TypeError, ValueError) as exc:
        raise error(400, "毕业年份格式不正确。") from exc
    if graduation_year and not 2000 <= graduation_year <= 2100:
        raise error(400, "毕业年份应在 2000 到 2100 之间。")
    invite_code = str(body.get("inviteCode") or f"CAMPUS-{uuid4().hex[:8]}").strip().upper()[:40]
    if one("SELECT id FROM campus_classes WHERE invite_code = ?", (invite_code,)):
        raise error(409, "班级邀请码已存在。")
    class_id = str(uuid4())
    db.execute(
        """
        INSERT INTO campus_classes (id, program_id, name, graduation_year, advisor, invite_code, status)
        VALUES (?, ?, ?, ?, ?, ?, 'active')
        """,
        (class_id, program_id, name, graduation_year, advisor or None, invite_code),
    )
    db.commit()
    record_audit(request, admin, "campus.class.create", target_type="campus_class", target_id=class_id, summary=f"创建班级：{name}")
    return {"class": one("SELECT * FROM campus_classes WHERE id = ?", (class_id,))}


def _delete_campus_enrollments_for_classes(class_ids: list[str]) -> None:
    """删除指定班级下学生的归班记录。"""
    if not class_ids:
        return
    placeholders = ",".join("?" for _ in class_ids)
    db.execute(f"DELETE FROM student_enrollments WHERE class_id IN ({placeholders})", class_ids)


@app.put("/api/admin/campus/colleges/{college_id}")
async def update_campus_college(request: Request, college_id: str, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    if not one("SELECT id FROM campus_colleges WHERE id = ?", (college_id,)):
        raise error(404, "学院不存在。")
    body = await request.json()
    code = _required_text(body, "code", "学院编码", 64).upper()
    name = _required_text(body, "name", "学院名称")
    dup = one("SELECT id FROM campus_colleges WHERE (code = ? OR name = ?) AND id != ?", (code, name, college_id))
    if dup:
        raise error(409, "学院编码或名称已存在。")
    db.execute("UPDATE campus_colleges SET code = ?, name = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (code, name, college_id))
    db.commit()
    record_audit(request, admin, "campus.college.update", target_type="campus_college", target_id=college_id, summary=f"更新学院：{name}")
    return {"college": one("SELECT * FROM campus_colleges WHERE id = ?", (college_id,))}


@app.delete("/api/admin/campus/colleges/{college_id}")
def delete_campus_college(request: Request, college_id: str, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    college = one("SELECT id, name FROM campus_colleges WHERE id = ?", (college_id,))
    if not college:
        raise error(404, "学院不存在。")
    class_rows = all_rows(
        """
        SELECT classes.id FROM campus_classes AS classes
        JOIN campus_programs AS programs ON programs.id = classes.program_id
        WHERE programs.college_id = ?
        """,
        (college_id,),
    )
    try:
        db.begin()
        _delete_campus_enrollments_for_classes([row["id"] for row in class_rows])
        db.execute("DELETE FROM campus_colleges WHERE id = ?", (college_id,))  # 级联删除其下专业、班级
        db.commit()
    except Exception:
        db.rollback()
        raise
    record_audit(request, admin, "campus.college.delete", target_type="campus_college", target_id=college_id, summary=f"级联删除学院：{college['name']}")
    return {"ok": True}


@app.put("/api/admin/campus/programs/{program_id}")
async def update_campus_program(request: Request, program_id: str, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    if not one("SELECT id FROM campus_programs WHERE id = ?", (program_id,)):
        raise error(404, "专业不存在。")
    body = await request.json()
    name = _required_text(body, "name", "专业名称")
    direction = str(body.get("direction") or "").strip()[:160]
    coordinator = str(body.get("coordinator") or "").strip()[:120]
    standard_major_code = str(body.get("standardMajorCode") or "").strip()[:64]
    db.execute(
        """
        UPDATE campus_programs
        SET standard_major_code = ?, name = ?, direction = ?, coordinator = ?, updated_at = CURRENT_TIMESTAMP
        WHERE id = ?
        """,
        (standard_major_code or None, name, direction or None, coordinator or None, program_id),
    )
    db.commit()
    record_audit(request, admin, "campus.program.update", target_type="campus_program", target_id=program_id, summary=f"更新专业：{name}")
    return {"program": one("SELECT * FROM campus_programs WHERE id = ?", (program_id,))}


@app.delete("/api/admin/campus/programs/{program_id}")
def delete_campus_program(request: Request, program_id: str, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    program = one("SELECT id, name FROM campus_programs WHERE id = ?", (program_id,))
    if not program:
        raise error(404, "专业不存在。")
    class_rows = all_rows("SELECT id FROM campus_classes WHERE program_id = ?", (program_id,))
    try:
        db.begin()
        _delete_campus_enrollments_for_classes([row["id"] for row in class_rows])
        db.execute("DELETE FROM campus_programs WHERE id = ?", (program_id,))  # 级联删除其下班级
        db.commit()
    except Exception:
        db.rollback()
        raise
    record_audit(request, admin, "campus.program.delete", target_type="campus_program", target_id=program_id, summary=f"级联删除专业：{program['name']}")
    return {"ok": True}


@app.put("/api/admin/campus/classes/{class_id}")
async def update_campus_class(request: Request, class_id: str, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    if not one("SELECT id FROM campus_classes WHERE id = ?", (class_id,)):
        raise error(404, "班级不存在。")
    body = await request.json()
    name = _required_text(body, "name", "班级名称")
    advisor = str(body.get("advisor") or "").strip()[:120]
    raw_year = body.get("graduationYear")
    try:
        graduation_year = int(raw_year) if raw_year not in {None, ""} else None
    except (TypeError, ValueError) as exc:
        raise error(400, "毕业年份格式不正确。") from exc
    if graduation_year and not 2000 <= graduation_year <= 2100:
        raise error(400, "毕业年份应在 2000 到 2100 之间。")
    invite_code = str(body.get("inviteCode") or "").strip().upper()[:40]
    if invite_code:
        dup = one("SELECT id FROM campus_classes WHERE invite_code = ? AND id != ?", (invite_code, class_id))
        if dup:
            raise error(409, "班级邀请码已存在。")
    db.execute(
        "UPDATE campus_classes SET name = ?, graduation_year = ?, advisor = ?, invite_code = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        (name, graduation_year, advisor or None, invite_code or None, class_id),
    )
    db.commit()
    record_audit(request, admin, "campus.class.update", target_type="campus_class", target_id=class_id, summary=f"更新班级：{name}")
    return {"class": one("SELECT * FROM campus_classes WHERE id = ?", (class_id,))}


@app.delete("/api/admin/campus/classes/{class_id}")
def delete_campus_class(request: Request, class_id: str, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    class_item = one("SELECT id, name FROM campus_classes WHERE id = ?", (class_id,))
    if not class_item:
        raise error(404, "班级不存在。")
    try:
        db.begin()
        _delete_campus_enrollments_for_classes([class_id])
        db.execute("DELETE FROM campus_classes WHERE id = ?", (class_id,))
        db.commit()
    except Exception:
        db.rollback()
        raise
    record_audit(request, admin, "campus.class.delete", target_type="campus_class", target_id=class_id, summary=f"删除班级：{class_item['name']}")
    return {"ok": True}


@app.put("/api/admin/campus/programs/{program_id}/jobs")
async def replace_program_jobs(request: Request, program_id: str, admin: dict = Depends(require_permission("manageOrganization"))):
    validate_admin_origin(request)
    if not one("SELECT id FROM campus_programs WHERE id = ?", (program_id,)):
        raise error(404, "专业不存在。")
    body = await request.json()
    job_role_ids = list(dict.fromkeys(str(item).strip() for item in (body.get("jobRoleIds") or []) if str(item).strip()))[:50]
    for job_role_id in job_role_ids:
        job = one(
            """
            SELECT jobs.id FROM catalog_job_roles AS jobs
            JOIN catalog_versions AS versions ON versions.id = jobs.version_id
            WHERE jobs.id = ? AND jobs.enabled = 1 AND versions.status = 'published'
            """,
            (job_role_id,),
        )
        if not job:
            raise error(400, "包含不存在或尚未发布的目标岗位。")
    try:
        db.begin()
        db.execute("DELETE FROM program_job_roles WHERE program_id = ?", (program_id,))
        for index, job_role_id in enumerate(job_role_ids):
            db.execute(
                "INSERT INTO program_job_roles (id, program_id, job_role_id, priority) VALUES (?, ?, ?, ?)",
                (str(uuid4()), program_id, job_role_id, "recommended" if index < 3 else "optional"),
            )
        db.commit()
    except Exception:
        db.rollback()
        raise
    record_audit(request, admin, "campus.program.jobs.update", target_type="campus_program", target_id=program_id, summary=f"更新推荐岗位：{len(job_role_ids)} 个")
    return {"ok": True, "jobRoleIds": job_role_ids}


@app.patch("/api/admin/campus/students/{user_id}")
async def update_campus_student(request: Request, user_id: str, admin: dict = Depends(require_permission("manageStudents"))):
    validate_admin_origin(request)
    _ensure_student_in_scope(admin, user_id)
    body = await request.json()
    existing = one("SELECT * FROM student_enrollments WHERE user_id = ?", (user_id,)) or {}
    class_id = body.get("classId", existing.get("class_id"))
    if class_id == "":
        class_id = None
    if class_id:
        _ensure_class_in_scope(admin, class_id)
    status = str(body.get("status", existing.get("status") or "active"))
    if status not in {"active", "inactive"}:
        raise error(400, "学生归属状态不正确。")
    enrollment = upsert_student_enrollment(
        user_id,
        class_id=class_id,
        student_no=str(body.get("studentNo", existing.get("student_no") or "")).strip(),
        status=status,
        focus_flag=bool(body.get("focus", existing.get("focus_flag") or False)),
        note=str(body.get("note", existing.get("note") or "")).strip(),
    )
    record_audit(request, admin, "campus.student.update", target_type="student", target_id=user_id, summary="更新学生组织归属与成长关注状态")
    return {"enrollment": enrollment}


@app.post("/api/admin/campus/students/import")
async def import_campus_students(
    request: Request,
    class_id: str,
    file: UploadFile = File(...),
    admin: dict = Depends(require_permission("manageStudents")),
):
    validate_admin_origin(request)
    if not one("SELECT id FROM campus_classes WHERE id = ?", (class_id,)):
        raise error(404, "目标班级不存在。")
    _ensure_class_in_scope(admin, class_id)
    content = await file.read()
    if len(content) > 2 * 1024 * 1024:
        raise error(400, "导入文件不能超过 2MB。")
    try:
        decoded = content.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(decoded)))
    except Exception as exc:
        raise error(400, "请上传 UTF-8 编码的 CSV 文件。") from exc
    if len(rows) > 1000:
        raise error(400, "单次最多导入 1000 名学生。")
    matched = 0
    unmatched: list[dict[str, str]] = []
    for index, row in enumerate(rows, start=2):
        email = normalize_email(row.get("邮箱") or row.get("email"))
        student_no = str(row.get("学号") or row.get("student_no") or "").strip()
        if not email:
            unmatched.append({"row": str(index), "email": "", "reason": "缺少邮箱"})
            continue
        user = one("SELECT id FROM users WHERE email = ?", (email,))
        if not user:
            unmatched.append({"row": str(index), "email": email, "reason": "学生尚未注册"})
            continue
        _ensure_student_in_scope(admin, user["id"])
        existing = one("SELECT * FROM student_enrollments WHERE user_id = ?", (user["id"],)) or {}
        upsert_student_enrollment(
            user["id"],
            class_id=class_id,
            student_no=student_no or str(existing.get("student_no") or ""),
            status=str(existing.get("status") or "active"),
            focus_flag=bool(existing.get("focus_flag") or False),
            note=str(existing.get("note") or ""),
        )
        matched += 1
    record_audit(request, admin, "campus.students.import", target_type="campus_class", target_id=class_id, summary=f"批量归班成功 {matched} 人，未匹配 {len(unmatched)} 人")
    return {"matched": matched, "unmatched": unmatched[:100], "total": len(rows)}


# ---------- 用户注册（学生登录白名单） ----------

@app.post("/api/admin/student-registrations/import")
async def import_student_registrations(
    _admin: dict = Depends(require_permission("manageStudents")),
):
    raise error(410, "旧学生注册入口已停用，请使用学生账号名单导入。")


@app.get("/api/admin/student-registrations")
def list_student_registrations(_admin: dict = Depends(require_permission("manageStudents"))):
    raise error(410, "旧学生注册入口已停用，请使用学生账号名单导入。")


@app.delete("/api/admin/student-registrations/{registration_id}")
def delete_student_registration(
    registration_id: str,
    _admin: dict = Depends(require_permission("manageStudents")),
):
    raise error(410, "旧学生注册入口已停用，历史数据不会删除。")
# ---------- 组织（组织结构一览 + 一键生成） ----------

@app.get("/api/admin/organization/structure")
def organization_structure(_admin: dict = Depends(require_admin)):
    """组织结构一览：学院 → 专业 → 班级（含学生数统计）。供组织管理与权限范围选择使用。"""
    data = campus_overview_data()
    return {
        "colleges": data["colleges"],
        "programs": data["programs"],
        "classes": data["classes"],
        "standardMajors": data["standardMajors"],
        "summary": data["summary"],
    }


@app.post("/api/admin/organization/import")
async def import_organization_structure(
    _admin: dict = Depends(require_permission("manageOrganization")),
):
    raise error(410, "旧组织导入会创建不可登录的注册记录，已停用。请先导入学生账号，再维护组织归属。")
@app.get("/api/admin/resume-form-settings")
def get_resume_form_settings(admin: dict = Depends(require_admin)):
    """读取候选人“填写简历”的学院/专业下拉配置。"""
    row = one("SELECT colleges FROM resume_form_settings ORDER BY updated_at DESC LIMIT 1")
    colleges = []
    if row:
        try:
            colleges = json.loads(row["colleges"] or "[]")
        except (TypeError, ValueError):
            colleges = []
    return {"colleges": colleges if isinstance(colleges, list) else []}


@app.put("/api/admin/resume-form-settings")
async def update_resume_form_settings(request: Request, admin: dict = Depends(require_super_admin)):
    """保存候选人“填写简历”的学院/专业下拉配置。"""
    validate_admin_origin(request)
    body = await request.json()
    colleges = body.get("colleges")
    if not isinstance(colleges, list):
        raise error(400, "学院配置格式不正确。")
    if len(colleges) > 200:
        raise error(400, "学院数量过多，请控制在 200 个以内。")
    for college in colleges:
        if not isinstance(college, dict) or not str(college.get("name") or "").strip():
            raise error(400, "每个学院都必须有名称。")
        majors = college.get("majors")
        if majors is not None and not isinstance(majors, list):
            raise error(400, "学院下的专业配置格式不正确。")
        if majors and len(majors) > 500:
            raise error(400, "单个学院的专业数量过多，请控制在 500 个以内。")
        for major in majors or []:
            if not isinstance(major, dict) or not str(major.get("name") or "").strip():
                raise error(400, "每个专业都必须有名称。")

    try:
        db.begin()
        db.execute("DELETE FROM resume_form_settings")
        db.execute(
            "INSERT INTO resume_form_settings (id, colleges, updated_by) VALUES (?, ?, ?)",
            (str(uuid4()), json.dumps(colleges, ensure_ascii=False), str(admin.get("email") or "")),
        )
        db.commit()
    except Exception:
        db.rollback()
        raise
    college_count = len(colleges)
    major_count = sum(len(college.get("majors") or []) for college in colleges)
    record_audit(
        request,
        admin,
        "resume_form_settings.update",
        target_type="resume_form_settings",
        summary=f"更新填写简历设置：{college_count} 个学院，{major_count} 个专业",
    )
    return {"ok": True, "colleges": colleges}


@app.get("/api/admin/health")
def health():
    return {"ok": True, "authRequired": True, "databasePath": get_database_path()}


# ---------------------------------------------------------------------------
# 招聘信息库：岗位管理 + 简历对比记录
# ---------------------------------------------------------------------------

JOB_POSTING_MAX_BYTES = 5 * 1024 * 1024

JOB_POSTING_EXPORT_COLUMNS = [
    ("title", "岗位名称"),
    ("company", "公司"),
    ("job_category", "岗位类别"),
    ("city", "工作地点"),
    ("graduation_year", "届次"),
    ("employment_type", "招聘类型"),
    ("salary", "薪资范围"),
    ("education_requirement", "学历要求"),
    ("experience_requirement", "经验要求"),
    ("headcount", "招聘人数"),
    ("deadline", "截止时间"),
    ("skills", "技能要求"),
    ("tags", "标签"),
    ("description", "岗位描述"),
    ("requirements", "任职要求"),
    ("catalog_job_name", "关联标准岗位"),
    ("status", "状态"),
    ("source_ref", "外部编号"),
]

JOB_POSTING_IMPORT_ALIASES = {
    "岗位名称": "title", "岗位": "title", "招聘岗位": "title", "职位名称": "title", "职位": "title",
    "公司": "company", "公司名称": "company", "企业名称": "company", "招聘公司": "company",
    "岗位类别": "job_category", "职位类别": "job_category", "岗位方向": "job_category",
    "工作地点": "city", "城市": "city", "地点": "city",
    "届次": "graduation_year", "毕业届次": "graduation_year", "招聘届次": "graduation_year",
    "招聘类型": "employment_type", "用工类型": "employment_type", "职位类型": "employment_type",
    "薪资": "salary", "薪资范围": "salary", "月薪": "salary", "待遇": "salary",
    "学历要求": "education_requirement", "学历": "education_requirement",
    "经验要求": "experience_requirement", "工作经验": "experience_requirement",
    "招聘人数": "headcount", "人数": "headcount",
    "截止时间": "deadline", "截止日期": "deadline", "招聘截止": "deadline", "有效期": "deadline",
    "技能要求": "skills", "技能": "skills", "技术栈": "skills",
    "标签": "tags",
    "岗位描述": "description", "岗位职责": "description", "职位描述": "description", "工作内容": "description",
    "任职要求": "requirements", "岗位要求": "requirements", "任职资格": "requirements",
    "关联标准岗位": "catalog_job_name", "标准岗位": "catalog_job_name",
    "状态": "status",
    "外部编号": "source_ref", "编号": "source_ref", "岗位编号": "source_ref",
}


def _published_job_role_map() -> dict[str, dict]:
    versions = list_versions(db, include_drafts=False)
    if not versions:
        return {}
    version_id = versions[0]["id"]
    rows = all_rows(
        "SELECT id, code, name FROM catalog_job_roles WHERE version_id = ? AND enabled = 1",
        (version_id,),
    )
    return {str(row["name"]).strip(): {"id": row["id"], "code": row["code"], "name": row["name"]} for row in rows}


def build_job_posting_workbook(rows: list[dict], *, template: bool = False) -> io.BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "招聘信息库"
    headers = [label for _field, label in JOB_POSTING_EXPORT_COLUMNS]
    worksheet.append(headers)
    header_fill = PatternFill("solid", fgColor="17324D")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        worksheet.append([
            (row.get("status") == "active" and "启用") or ("停用" if row.get("status") == "disabled" else (row.get("status") or ""))
            if field == "status"
            else (row.get(field) if row.get(field) is not None else "")
            for field, _label in JOB_POSTING_EXPORT_COLUMNS
        ])

    widths = [24, 20, 16, 14, 10, 12, 14, 12, 12, 10, 14, 26, 18, 40, 40, 18, 10, 16]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    if template:
        notes = workbook.create_sheet("填写说明")
        notes.append(["字段", "是否必填", "说明"])
        for cell in notes[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
        for field, label in JOB_POSTING_EXPORT_COLUMNS:
            required = "是" if field == "title" else "否"
            hint = {
                "title": "招聘岗位名称，必填。",
                "skills": "多个技能用中文逗号、或顿号分隔。",
                "description": "岗位职责/工作内容，可多行。",
                "requirements": "任职要求，建议每行一条，便于逐条对比。",
                "catalog_job_name": "可选，填写岗位知识库中已发布的标准岗位名称即可自动关联。",
                "status": "填“启用”或“停用”，默认启用。",
                "source_ref": "外部唯一编号；相同编号再次导入会更新原记录，留空则按“岗位名称+公司”去重。",
            }.get(field, "")
            if hint:
                notes.append([label, required, hint])
        notes.column_dimensions["A"].width = 18
        notes.column_dimensions["B"].width = 10
        notes.column_dimensions["C"].width = 80

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def evaluate_job_posting_import(content: bytes) -> dict[str, Any]:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise error(400, f"无法解析 Excel 文件：{exc}") from exc

    worksheet = workbook.worksheets[0] if workbook.worksheets else None
    if worksheet is None:
        raise error(400, "Excel 中没有可读取的工作表。")

    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        raise error(400, "Excel 内容为空。")

    field_by_index: dict[int, str] = {}
    for index, cell in enumerate(header_row or []):
        label = str(cell or "").strip()
        field = JOB_POSTING_IMPORT_ALIASES.get(label)
        if field:
            field_by_index[index] = field
    if "title" not in field_by_index.values():
        raise error(400, "Excel 缺少“岗位名称”列，请参考导入模板。")

    job_roles = _published_job_role_map()
    result = {"total": 0, "created": 0, "updated": 0, "skipped": 0, "errors": []}
    for row_number, raw_row in enumerate(rows, start=2):
        if raw_row is None:
            continue
        payload: dict[str, Any] = {}
        for index, field in field_by_index.items():
            if index >= len(raw_row):
                continue
            value = raw_row[index]
            if value is None:
                continue
            payload[field] = value if field == "headcount" else str(value).strip()
        if not str(payload.get("title") or "").strip():
            result["skipped"] += 1
            continue

        result["total"] += 1
        catalog_name = str(payload.pop("catalog_job_name", "") or "").strip()
        if catalog_name:
            matched = job_roles.get(catalog_name)
            if matched:
                payload["catalog_job_role_id"] = matched["id"]
                payload["catalog_job_name"] = matched["name"]
            else:
                payload["catalog_job_role_id"] = None
                payload["catalog_job_name"] = catalog_name

        status_label = str(payload.get("status") or "").strip()
        if status_label:
            payload["status"] = "active" if status_label in {"启用", "正常", "active", "有效"} else "disabled"

        source_ref = str(payload.get("source_ref") or "").strip()
        if not source_ref:
            source_ref = f"{payload.get('title')}|{payload.get('company') or ''}"
        payload["source_ref"] = source_ref

        try:
            _posting, action = upsert_job_posting_by_source(
                db,
                payload,
                source="excel",
                source_ref=source_ref,
            )
            result["created" if action == "created" else "updated"] += 1
        except ValueError as exc:
            result["skipped"] += 1
            if len(result["errors"]) < 20:
                result["errors"].append(f"第 {row_number} 行：{exc}")
    db.commit()
    return result


def _job_posting_scope_note() -> str:
    return "招聘信息库包含手动录入与 Excel 导入的岗位，可选择性关联岗位知识库中的标准岗位。"


@app.get("/api/admin/job-postings")
def admin_list_job_postings(
    keyword: str = "",
    status: str = "",
    admin: dict = Depends(require_permission("manageCatalog")),
):
    postings = list_job_postings(db, keyword=keyword.strip() or None, status=status.strip() or None)
    return {
        "postings": postings,
        "jobRoles": [
            {"id": role["id"], "name": role["name"]}
            for role in _published_job_role_map().values()
        ],
        "scopeNote": _job_posting_scope_note(),
    }


@app.get("/api/admin/job-postings/template")
def admin_download_job_posting_template(
    request: Request,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    output = build_job_posting_workbook([], template=True)
    record_audit(request, admin, "job_postings.template", target_type="job_posting", summary="下载招聘信息导入模板")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="job-posting-import-template.xlsx"'},
    )


@app.get("/api/admin/job-postings/export")
def admin_export_job_postings(
    request: Request,
    keyword: str = "",
    status: str = "",
    admin: dict = Depends(require_permission("manageCatalog")),
):
    postings = list_job_postings(db, keyword=keyword.strip() or None, status=status.strip() or None)
    output = build_job_posting_workbook(postings)
    record_audit(
        request,
        admin,
        "job_postings.export",
        target_type="job_posting",
        summary=f"导出 {len(postings)} 条招聘岗位",
    )
    filename = f"job-postings-{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/admin/job-postings/import")
async def admin_import_job_postings(
    request: Request,
    file: UploadFile = File(...),
    admin: dict = Depends(require_permission("manageCatalog")),
):
    validate_admin_origin(request)
    filename = str(file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise error(400, "请上传 .xlsx 格式的招聘岗位表。")
    content = await file.read(JOB_POSTING_MAX_BYTES + 1)
    if not content:
        raise error(400, "上传的文件为空。")
    if len(content) > JOB_POSTING_MAX_BYTES:
        raise error(413, "招聘岗位表不能超过 5MB。")

    result = evaluate_job_posting_import(content)
    record_audit(
        request,
        admin,
        "job_postings.import",
        target_type="job_posting",
        summary=f"导入招聘岗位：新增 {result['created']} 条，更新 {result['updated']} 条，跳过 {result['skipped']} 条",
    )
    return {"ok": True, **result}


@app.post("/api/admin/job-postings/sync")
async def admin_sync_job_postings(
    request: Request,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    """预留的第三方招聘数据源同步入口。"""
    validate_admin_origin(request)
    endpoint = str(os.environ.get("JOB_POSTING_SYNC_ENDPOINT") or "").strip()
    if not endpoint:
        raise error(
            400,
            "尚未配置第三方招聘数据源。请设置环境变量 JOB_POSTING_SYNC_ENDPOINT 后再使用同步功能。",
        )
    raise error(
        501,
        "第三方招聘数据源同步接口已预留，但当前版本尚未接入具体数据源（"
        f"{endpoint}）。请先使用手动录入或 Excel 导入。",
    )


@app.post("/api/admin/job-postings", status_code=201)
async def admin_create_job_posting(
    request: Request,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    validate_admin_origin(request)
    body = await request.json()
    if not isinstance(body, dict):
        raise error(400, "请求格式不正确。")
    payload = dict(body)
    catalog_name = str(payload.get("catalog_job_name") or "").strip()
    if catalog_name:
        matched = _published_job_role_map().get(catalog_name)
        if matched:
            payload["catalog_job_role_id"] = matched["id"]
            payload["catalog_job_name"] = matched["name"]
    try:
        posting = create_job_posting(db, payload, created_by=str(admin.get("email") or ""))
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    db.commit()
    record_audit(
        request,
        admin,
        "job_postings.create",
        target_type="job_posting",
        target_id=posting.get("id"),
        summary=f"新增招聘岗位：{posting.get('title')}",
    )
    return {"ok": True, "posting": posting}


@app.put("/api/admin/job-postings/{posting_id}")
async def admin_update_job_posting(
    request: Request,
    posting_id: str,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    validate_admin_origin(request)
    body = await request.json()
    if not isinstance(body, dict):
        raise error(400, "请求格式不正确。")
    payload = dict(body)
    catalog_name = str(payload.get("catalog_job_name") or "").strip()
    if catalog_name:
        matched = _published_job_role_map().get(catalog_name)
        payload["catalog_job_role_id"] = matched["id"] if matched else None
        if matched:
            payload["catalog_job_name"] = matched["name"]
    try:
        posting = update_job_posting(db, posting_id, payload)
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    if not posting:
        raise error(404, "招聘岗位不存在或已被删除。")
    db.commit()
    record_audit(
        request,
        admin,
        "job_postings.update",
        target_type="job_posting",
        target_id=posting_id,
        summary=f"更新招聘岗位：{posting.get('title')}",
    )
    return {"ok": True, "posting": posting}


@app.delete("/api/admin/job-postings/{posting_id}")
def admin_delete_job_posting(
    request: Request,
    posting_id: str,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    validate_admin_origin(request)
    posting = get_job_posting(db, posting_id)
    if not posting:
        raise error(404, "招聘岗位不存在或已被删除。")
    delete_job_posting(db, posting_id)
    db.commit()
    record_audit(
        request,
        admin,
        "job_postings.delete",
        target_type="job_posting",
        target_id=posting_id,
        summary=f"删除招聘岗位：{posting.get('title')}",
    )
    return {"ok": True}


@app.get("/api/admin/job-matches")
def admin_list_job_matches(
    keyword: str = "",
    jobPostingId: str = "",
    limit: int = Query(200, ge=1, le=1000),
    admin: dict = Depends(require_permission("manageCatalog")),
):
    matches = list_job_matches(
        db,
        keyword=keyword.strip() or None,
        job_posting_id=jobPostingId.strip() or None,
        limit=limit,
    )
    return {
        "matches": matches,
        "postings": [
            {"id": posting["id"], "title": posting["title"], "company": posting.get("company")}
            for posting in list_job_postings(db, limit=1000)
        ],
        "scopeNote": "对比记录仅保存简历摘要与匹配结论，用于就业指导，不包含完整简历正文。",
    }


@app.get("/api/admin/job-matches/{match_id}")
def admin_get_job_match(
    match_id: str,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    match = get_job_match(db, match_id)
    if not match:
        raise error(404, "对比记录不存在。")
    return {"match": match}


def _jd_submission_scope_note() -> str:
    return "学生粘贴的岗位描述（JD）在此统一管理，审核通过后会写入招聘信息库的岗位管理列表。"


@app.get("/api/admin/jd-submissions")
def admin_list_jd_submissions(
    keyword: str = "",
    status: str = "",
    limit: int = Query(200, ge=1, le=1000),
    admin: dict = Depends(require_permission("manageCatalog")),
):
    submissions = list_jd_submissions(
        db,
        keyword=keyword.strip() or None,
        status=status.strip() or None,
        limit=limit,
    )
    return {"submissions": submissions, "scopeNote": _jd_submission_scope_note()}


@app.post("/api/admin/jd-submissions", status_code=201)
async def admin_create_jd_submission(
    request: Request,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    validate_admin_origin(request)
    body = await request.json()
    if not isinstance(body, dict):
        raise error(400, "请求格式不正确。")
    payload = dict(body)
    title = str(payload.get("job_title") or "").strip()
    if not title:
        raise error(400, "岗位名称不能为空。")
    submission = create_jd_submission(
        db,
        job_title=title,
        company=str(payload.get("company") or "").strip(),
        jd_text=str(payload.get("jd_text") or "").strip(),
        student_name=str(payload.get("student_name") or "管理员录入").strip(),
    )
    db.commit()
    record_audit(
        request,
        admin,
        "jd_submissions.create",
        target_type="jd_submission",
        target_id=submission.get("id"),
        summary=f"新增粘贴 JD 记录：{submission.get('job_title')}",
    )
    return {"ok": True, "submission": submission}


@app.put("/api/admin/jd-submissions/{submission_id}")
async def admin_update_jd_submission(
    request: Request,
    submission_id: str,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    validate_admin_origin(request)
    body = await request.json()
    if not isinstance(body, dict):
        raise error(400, "请求格式不正确。")
    submission = update_jd_submission(
        db,
        submission_id,
        dict(body),
        decided_by=str(admin.get("email") or ""),
    )
    if not submission:
        raise error(404, "粘贴 JD 记录不存在或已被删除。")
    db.commit()
    record_audit(
        request,
        admin,
        "jd_submissions.update",
        target_type="jd_submission",
        target_id=submission_id,
        summary=f"更新粘贴 JD 记录：{submission.get('job_title')}",
    )
    return {"ok": True, "submission": submission}


@app.delete("/api/admin/jd-submissions/{submission_id}")
def admin_delete_jd_submission(
    request: Request,
    submission_id: str,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    validate_admin_origin(request)
    submission = get_jd_submission(db, submission_id)
    if not submission:
        raise error(404, "粘贴 JD 记录不存在或已被删除。")
    delete_jd_submission(db, submission_id)
    db.commit()
    record_audit(
        request,
        admin,
        "jd_submissions.delete",
        target_type="jd_submission",
        target_id=submission_id,
        summary=f"删除粘贴 JD 记录：{submission.get('job_title')}",
    )
    return {"ok": True}


@app.post("/api/admin/jd-submissions/{submission_id}/approve")
async def admin_approve_jd_submission(
    request: Request,
    submission_id: str,
    admin: dict = Depends(require_permission("manageCatalog")),
):
    validate_admin_origin(request)
    submission = get_jd_submission(db, submission_id)
    if not submission:
        raise error(404, "粘贴 JD 记录不存在或已被删除。")
    try:
        body = await request.json()
    except Exception:
        body = None
    payload = dict(body) if isinstance(body, dict) else {}
    title = str(payload.get("job_title") or submission.get("job_title") or "").strip() or "自定义岗位"
    company = str(payload.get("company") or submission.get("company") or "").strip()
    jd_text = str(payload.get("jd_text") or submission.get("jd_text") or "").strip()

    updated = update_jd_submission(
        db,
        submission_id,
        {"job_title": title, "company": company, "jd_text": jd_text},
        decided_by=str(admin.get("email") or ""),
    )
    if not updated:
        raise error(404, "粘贴 JD 记录不存在或已被删除。")

    posting, action = upsert_job_posting_by_source(
        db,
        {
            "title": title,
            "company": company,
            "description": jd_text,
            "requirements": jd_text,
            "status": "active",
            "published": 1,
        },
        source="student_jd",
        source_ref=submission_id,
        created_by=str(admin.get("email") or ""),
    )
    submission = mark_jd_submission_approved(
        db,
        submission_id,
        str(posting.get("id") or ""),
        decided_by=str(admin.get("email") or ""),
    )
    db.commit()
    record_audit(
        request,
        admin,
        "jd_submissions.approve",
        target_type="jd_submission",
        target_id=submission_id,
        summary=f"粘贴 JD 审核入库：{title}",
    )
    return {"ok": True, "submission": submission, "posting": posting, "action": action}


@app.get("/api/admin/catalog/versions")
def admin_catalog_versions(_admin: dict = Depends(require_catalog_permission("read"))):
    return {"versions": list_versions(db, include_drafts=True)}


@app.post("/api/admin/catalog/versions", status_code=201)
async def admin_create_catalog_version(request: Request, admin: dict = Depends(require_catalog_permission("write"))):
    validate_admin_origin(request)
    try:
        version = create_version(db, await request.json(), admin.get("id"))
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    record_audit(request, admin, "catalog.version.create", target_type="catalog_version", target_id=version["id"], summary=version["code"])
    return {"version": version}


@app.get("/api/admin/catalog/tree")
def admin_catalog_tree(version: str | None = None, _admin: dict = Depends(require_catalog_permission("read"))):
    result = catalog_tree(db, version_code=version, include_disabled=True, include_drafts=True)
    if not result:
        raise error(404, "职业能力目录版本不存在。")
    return result


@app.post("/api/admin/catalog/versions/{version_id}/entities/{entity_type}", status_code=201)
async def admin_create_catalog_entity(
    request: Request,
    version_id: str,
    entity_type: str,
    admin: dict = Depends(require_catalog_permission("write")),
):
    validate_admin_origin(request)
    try:
        item = create_catalog_entity(db, version_id, entity_type, await request.json())
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    record_audit(
        request,
        admin,
        f"catalog.{entity_type}.create",
        target_type=entity_type,
        target_id=item["id"],
        summary=item.get("name") or item.get("code"),
    )
    return {"item": item}


@app.patch("/api/admin/catalog/versions/{version_id}/entities/{entity_type}/{entity_id}")
async def admin_update_catalog_entity(
    request: Request,
    version_id: str,
    entity_type: str,
    entity_id: str,
    admin: dict = Depends(require_catalog_permission("write")),
):
    validate_admin_origin(request)
    try:
        item = update_catalog_entity(db, version_id, entity_type, entity_id, await request.json())
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    if not item:
        raise error(404, "目录数据不存在。")
    record_audit(
        request,
        admin,
        f"catalog.{entity_type}.update",
        target_type=entity_type,
        target_id=entity_id,
        summary=item.get("name") or item.get("code"),
    )
    return {"item": item}


@app.delete("/api/admin/catalog/versions/{version_id}/entities/{entity_type}/{entity_id}")
def admin_delete_catalog_entity(
    request: Request,
    version_id: str,
    entity_type: str,
    entity_id: str,
    admin: dict = Depends(require_catalog_permission("write")),
):
    validate_admin_origin(request)
    try:
        deleted = delete_catalog_entity(db, version_id, entity_type, entity_id)
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    if not deleted:
        raise error(404, "目录数据不存在。")
    record_audit(request, admin, f"catalog.{entity_type}.delete", target_type=entity_type, target_id=entity_id)
    return {"ok": True}


@app.put("/api/admin/catalog/versions/{version_id}/jobs/{job_role_id}/competencies")
async def admin_replace_job_competencies(
    request: Request,
    version_id: str,
    job_role_id: str,
    admin: dict = Depends(require_catalog_permission("write")),
):
    validate_admin_origin(request)
    body = await request.json()
    items = body.get("competencies")
    if not isinstance(items, list):
        raise error(400, "岗位能力配置必须是数组。")
    try:
        competencies = replace_job_competencies(db, version_id, job_role_id, items)
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    record_audit(
        request,
        admin,
        "catalog.job.competencies.replace",
        target_type="job_role",
        target_id=job_role_id,
        summary=f"{len(competencies)} 项能力要求",
    )
    return {"competencies": competencies}


@app.post("/api/admin/catalog/versions/{version_id}/publish")
def admin_publish_catalog_version(request: Request, version_id: str, admin: dict = Depends(require_catalog_permission("publish"))):
    validate_admin_origin(request)
    try:
        version = publish_version(db, version_id, admin.get("id"))
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    if not version:
        raise error(404, "职业能力目录版本不存在。")
    record_audit(request, admin, "catalog.version.publish", target_type="catalog_version", target_id=version_id, summary=version["code"])
    return {"version": version}


@app.post("/api/admin/catalog/versions/{version_id}/imports", status_code=201)
async def admin_import_catalog(
    request: Request,
    version_id: str,
    file: UploadFile = File(...),
    mode: str = "merge",
    admin: dict = Depends(require_catalog_permission("import")),
):
    validate_admin_origin(request)
    filename = (file.filename or "catalog.xlsx").strip()
    if not filename.lower().endswith(".xlsx"):
        raise error(400, "目录导入仅支持 .xlsx 文件。")
    content = await file.read()
    if not content or len(content) > 10 * 1024 * 1024:
        raise error(400, "Excel 文件不能为空且不能超过 10MB。")
    try:
        job = import_catalog_excel(db, version_id, filename, content, mode, admin.get("id"))
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    record_audit(request, admin, "catalog.excel.import", target_type="catalog_version", target_id=version_id, summary=f"{filename}，{job['imported_rows']} 行")
    return {"import_job": job}


@app.get("/api/admin/catalog/job-suggestions")
def admin_job_suggestions(status: str | None = None, _admin: dict = Depends(require_catalog_permission("read"))):
    if status and status not in {"pending", "approved", "rejected", "merged"}:
        raise error(400, "岗位建议状态不正确。")
    return {"suggestions": list_job_suggestions(db, status)}


@app.post("/api/admin/catalog/job-suggestions/{suggestion_id}/review")
async def admin_review_job_suggestion(
    request: Request,
    suggestion_id: str,
    admin: dict = Depends(require_catalog_permission("write")),
):
    validate_admin_origin(request)
    body = await request.json()
    status = str(body.get("status") or "").strip()
    try:
        suggestion = review_job_suggestion(db, suggestion_id, status, admin.get("id"), body.get("note"))
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    if not suggestion:
        raise error(404, "岗位建议不存在。")
    record_audit(request, admin, f"catalog.suggestion.{status}", target_type="job_suggestion", target_id=suggestion_id, summary=suggestion["suggested_name"])
    return {"suggestion": suggestion}


@app.post("/api/admin/catalog/job-suggestions/{suggestion_id}/merge")
async def admin_merge_job_suggestion(
    request: Request,
    suggestion_id: str,
    admin: dict = Depends(require_catalog_permission("write")),
):
    validate_admin_origin(request)
    body = await request.json()
    job_role_id = str(body.get("job_role_id") or "").strip()
    if not job_role_id:
        raise error(400, "请选择需要合并到的正式岗位。")
    try:
        suggestion = merge_job_suggestion(db, suggestion_id, job_role_id, admin.get("id"))
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    if not suggestion:
        raise error(404, "岗位建议不存在。")
    record_audit(request, admin, "catalog.suggestion.merge", target_type="job_suggestion", target_id=suggestion_id, summary=suggestion["suggested_name"])
    return {"suggestion": suggestion}


@app.post("/api/admin/catalog/job-suggestions/{suggestion_id}/approve-to-draft", status_code=201)
async def admin_approve_job_suggestion_to_draft(
    request: Request,
    suggestion_id: str,
    admin: dict = Depends(require_catalog_permission("write")),
):
    validate_admin_origin(request)
    body = await request.json()
    version_id = str(body.get("version_id") or "").strip()
    if not version_id:
        raise error(400, "请选择需要写入的目录草稿。")
    try:
        result = approve_job_suggestion_to_draft(db, suggestion_id, version_id, body, admin.get("id"))
    except ValueError as exc:
        raise error(400, str(exc)) from exc
    if not result:
        raise error(404, "岗位建议不存在。")
    record_audit(
        request,
        admin,
        "catalog.suggestion.approve_to_draft",
        target_type="job_suggestion",
        target_id=suggestion_id,
        summary=f"{result['suggestion']['suggested_name']} → {result['job_role']['code']}",
    )
    return result


@app.post("/api/admin/auth/login")
async def admin_login(request: Request, response: Response):
    validate_admin_origin(request)
    body = await request.json()
    email = normalize_email(body.get("email"))
    password = str(body.get("password") or "")
    if not is_valid_email(email) or not password:
        raise error(400, "请输入有效的管理员邮箱和密码。")
    if is_login_limited(request, email):
        raise error(429, "登录尝试过于频繁，请稍后再试。")
    if email == MASTER_ADMIN_EMAIL:
        bootstrap_master_admin()  # 代码定义的主管理员账号始终可用
    admin = one(
        """
        SELECT id, email, password_hash, name, role, status, permissions, student_scope, last_login_at
        FROM admin_users WHERE email = ?
        """,
        (email,),
    )
    if not admin or admin.get("status") != "normal" or not verify_password(password, admin.get("password_hash") or ""):
        record_failed_login(request, email)
        record_audit(
            request,
            {"email": email},
            "admin.login_failed",
            target_type="admin_user",
            target_id=email,
            summary="管理员登录失败",
            success=False,
        )
        raise error(401, "管理员邮箱或密码不正确。")
    clear_failed_logins(request, email)
    db.execute("UPDATE admin_users SET last_login_at = CURRENT_TIMESTAMP WHERE id = ?", (admin["id"],))
    db.commit()
    create_admin_session(response, request, admin["id"])
    current_admin = one(
        "SELECT id, email, name, role, status, permissions, student_scope, last_login_at FROM admin_users WHERE id = ?",
        (admin["id"],),
    )
    record_audit(request, current_admin, "admin.login", target_type="admin_user", target_id=admin["id"], summary="管理员登录成功")
    return {"admin": sanitize_admin(current_admin)}


@app.get("/api/admin/auth/me")
def admin_me(admin: dict = Depends(require_admin)):
    return {"admin": sanitize_admin(admin)}


@app.post("/api/admin/auth/logout")
def admin_logout(request: Request, response: Response, admin: dict = Depends(require_admin)):
    validate_admin_origin(request)
    token = request.cookies.get(ADMIN_SESSION_COOKIE)
    if token:
        db.execute("DELETE FROM admin_sessions WHERE token_hash = ?", (hash_token(token),))
        db.commit()
    record_audit(request, admin, "admin.logout", target_type="admin_user", target_id=admin["id"], summary="管理员退出登录")
    clear_admin_session(response)
    return {"ok": True}


@app.get("/api/admin/connection-logs")
def admin_connection_logs(
    level: str = "",
    event_type: str = "",
    query: str = "",
    limit: int = 300,
    _admin: dict = Depends(require_permission("viewConnectionLogs")),
):
    normalized_level = level.strip().lower()
    if normalized_level and normalized_level not in {"info", "warning", "error"}:
        raise error(400, "连接日志级别不正确。")
    normalized_event_type = event_type.strip()[:80]
    normalized_query = query.strip()[:100]
    return {
        "logs": list_connection_logs(
            level=normalized_level,
            event_type=normalized_event_type,
            query=normalized_query,
            limit=limit,
        ),
        "summary": connection_log_summary(),
        "retentionDays": WEBRTC_DIAGNOSTIC_RETENTION_DAYS,
    }


@app.get("/api/admin/snapshot")
def snapshot(admin: dict = Depends(require_admin)):
    try:
        reports = list_reports()
        is_super = is_super_admin(admin)
        perms = set(ALL_PERMISSIONS) if is_super else admin_permissions(admin)
        can_students = "manageStudents" in perms
        can_interviews = "viewInterviews" in perms
        can_reports = "viewReports" in perms
        can_model = "manageModelConfig" in perms
        can_catalog = "manageCatalog" in perms
        scope = admin_student_scope(admin)
        return {
            "metrics": build_metrics(reports),
            "candidates": list_candidates(scope=scope) if can_students else [],
            "interviews": list_interviews() if can_interviews else [],
            "reports": reports if can_reports else [],
            "agents": list_agent_templates() if can_model else [],
            "auditLogs": list_audit_logs(30) if is_super else [],
            "adminUsers": list_admin_accounts() if is_super else [],
            "settings": build_settings() if is_super else {},
            "permissions": {
                "canViewCandidates": can_students or can_interviews,
                "canViewInterviews": can_interviews,
                "canViewReports": can_reports,
                "canViewAgents": can_model,
                "canViewConnectionLogs": "viewConnectionLogs" in perms,
                "canManageStudents": can_students,
                "canImportStudentAccounts": is_super,
                "canManageOrganization": "manageOrganization" in perms,
                "canManageCampus": can_students or "manageOrganization" in perms,
                "canManageSettings": is_super,
                "canViewAudit": is_super,
                "canViewCatalog": can_catalog,
                "canWriteCatalog": can_catalog,
                "canImportCatalog": can_catalog,
                "canPublishCatalog": can_catalog,
                "canViewJobPostings": can_catalog or is_super,
            },
            "admin": sanitize_admin(admin),
        }
    except Exception as exc:
        logger.exception("Management snapshot failed")
        raise error(500, "管理端数据读取失败，请稍后重试。") from exc


@app.post("/api/admin/student-accounts/import")
async def import_student_accounts(
    request: Request,
    file: UploadFile = File(...),
    admin: dict = Depends(require_super_admin),
):
    validate_admin_origin(request)
    filename = str(file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise error(400, "请上传 .xlsx 格式的学生名单。")
    content = await file.read(STUDENT_IMPORT_MAX_BYTES + 1)
    if not content:
        raise error(400, "上传的学生名单为空。")
    if len(content) > STUDENT_IMPORT_MAX_BYTES:
        raise error(413, "学生名单不能超过 10MB。")

    result = import_student_accounts_from_workbook(content)
    record_audit(
        request,
        admin,
        "student_accounts.import",
        target_type="student_account",
        summary=f"导入学生名单：新增 {result['created']}，更新 {result['updated']}，迁移旧账号 {result['migrated']}，跳过 {result['skipped']}",
    )
    return {"ok": True, "result": result}


@app.get("/api/admin/student-accounts/export")
def export_student_accounts(
    request: Request,
    counselor: str = "",
    class_name: str = "",
    admission_year: str = "",
    admin: dict = Depends(require_permission("manageStudents")),
):
    normalized_counselor = counselor.strip()[:120]
    normalized_class_name = class_name.strip()[:160]
    normalized_admission_year = admission_year.strip()
    if normalized_admission_year and not re.match(r"^20\d{2}$", normalized_admission_year):
        raise error(400, "学年筛选格式不正确。")
    clauses = ["users.must_change_password = 1", "users.temp_password_encrypted IS NOT NULL"]
    params: list[Any] = []
    if normalized_counselor:
        clauses.append("users.counselor = ?")
        params.append(normalized_counselor)
    if normalized_class_name:
        clauses.append("users.class_name = ?")
        params.append(normalized_class_name)
    if normalized_admission_year:
        clauses.append("users.student_no LIKE ?")
        params.append(f"{normalized_admission_year}%")
    scope = admin_student_scope(admin)
    if scope is not None:
        clauses.append("classes.id IS NOT NULL")
        allowed_scopes = []
        for entry in scope:
            fields = []
            for column, key in (("programs.college_id", "college"), ("programs.id", "program"), ("classes.id", "class")):
                if entry.get(key):
                    fields.append(f"{column} = ?")
                    params.append(entry[key])
            allowed_scopes.append("(" + " AND ".join(fields or ["1=1"]) + ")")
        clauses.append("(" + " OR ".join(allowed_scopes or ["1=0"]) + ")")
    where = "WHERE " + " AND ".join(clauses)
    rows = all_rows(
        f"""
        SELECT users.student_no, users.name, users.college, users.gender,
               users.class_name, users.counselor, users.source_account_status,
               users.student_status, users.status, users.temp_password_encrypted,
               users.temp_password_created_at
        FROM users
        LEFT JOIN student_enrollments AS enrollments ON enrollments.user_id = users.id
        LEFT JOIN campus_classes AS classes ON classes.id = enrollments.class_id
        LEFT JOIN campus_programs AS programs ON programs.id = classes.program_id
        {where}
        ORDER BY users.counselor, users.college, users.class_name, users.student_no
        """,
        tuple(params),
    )
    output = build_student_password_workbook(rows)
    record_audit(
        request,
        admin,
        "student_accounts.export",
        target_type="student_account",
        target_id=normalized_counselor,
        summary=(
            f"导出 {len(rows)} 名待激活学生的临时密码"
            + (f"，辅导员：{normalized_counselor}" if normalized_counselor else "")
            + (f"，班级：{normalized_class_name}" if normalized_class_name else "")
            + (f"，学年：{normalized_admission_year}" if normalized_admission_year else "")
        ),
    )
    filename = f"student-temporary-passwords-{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/admin/student-accounts/{user_id}/temporary-password")
def reveal_student_temporary_password(
    request: Request,
    user_id: str,
    admin: dict = Depends(require_permission("manageStudents")),
):
    user = one(
        """
        SELECT id, student_no, must_change_password, temp_password_encrypted
        FROM users WHERE id = ?
        """,
        (user_id,),
    )
    if not user:
        raise error(404, "学生账号不存在。")
    _ensure_student_in_scope(admin, user_id)
    if not user.get("must_change_password") or not user.get("temp_password_encrypted"):
        raise error(409, "该学生已完成改密，临时密码不可查看。")
    try:
        temporary_password = decrypt_temporary_password(user["temp_password_encrypted"])
    except ValueError as exc:
        raise error(409, str(exc)) from exc
    record_audit(
        request,
        admin,
        "student_accounts.password_reveal",
        target_type="student_account",
        target_id=user_id,
        summary=f"查看学号 {user.get('student_no') or '-'} 的临时密码",
    )
    return {"temporaryPassword": temporary_password}


@app.post("/api/admin/student-accounts/{user_id}/reset-password")
def reset_student_password(
    request: Request,
    user_id: str,
    admin: dict = Depends(require_permission("manageStudents")),
):
    validate_admin_origin(request)
    user = one("SELECT id, student_no FROM users WHERE id = ?", (user_id,))
    if not user or not user.get("student_no"):
        raise error(404, "学生账号不存在或尚未绑定学号。")
    _ensure_student_in_scope(admin, user_id)
    temporary_password = generate_temporary_password()
    db.begin()
    try:
        cursor = db.execute(
            """
            UPDATE users
            SET password_hash = ?, must_change_password = 1, temp_password_encrypted = ?,
                temp_password_created_at = CURRENT_TIMESTAMP, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (hash_password(temporary_password), encrypt_temporary_password(temporary_password), user_id),
        )
        cursor.close()
        cursor = db.execute("DELETE FROM sessions WHERE user_id = ?", (user_id,))
        cursor.close()
        db.commit()
    except Exception:
        db.rollback()
        raise
    record_audit(
        request,
        admin,
        "student_accounts.password_reset",
        target_type="student_account",
        target_id=user_id,
        summary=f"重置学号 {user['student_no']} 的临时密码并撤销全部登录会话",
    )
    return {"temporaryPassword": temporary_password, "mustChangePassword": True}


@app.get("/api/admin/candidates/{candidate_id}")
def candidate_detail(request: Request, candidate_id: str, admin: dict = Depends(require_permission("manageStudents"))):
    _ensure_student_in_scope(admin, candidate_id)
    result = {"candidate": get_candidate_detail(candidate_id)}
    record_audit(request, admin, "candidate.view", target_type="candidate", target_id=candidate_id, summary="查看候选人详情")
    return result


@app.get("/api/admin/interviews/{interview_id}")
def interview_detail(request: Request, interview_id: str, admin: dict = Depends(require_permission("viewInterviews"))):
    result = get_interview_detail(interview_id)
    record_audit(request, admin, "interview.view", target_type="interview", target_id=interview_id, summary="查看面试详情")
    return result


@app.get("/api/admin/reports/{report_id}")
def report_detail(request: Request, report_id: str, admin: dict = Depends(require_permission("viewReports"))):
    result = {"report": get_report_detail(report_id)}
    record_audit(request, admin, "report.view", target_type="report", target_id=report_id, summary="查看面试报告")
    return result


@app.patch("/api/admin/reports/{report_id}/review")
async def review_report(request: Request, report_id: str, admin: dict = Depends(require_permission("viewReports"))):
    validate_admin_origin(request)
    body = await request.json()
    status = str(body.get("status") or "").strip()
    if status not in {"approved", "rejected", "pending"}:
        raise error(400, "复核状态必须是 approved、rejected 或 pending。")
    if not one("SELECT id FROM interview_reports WHERE id = ?", (report_id,)):
        raise error(404, "报告不存在。")
    db.execute(
        """UPDATE interview_reports
           SET review_status = ?, reviewed_by = ?, reviewed_at = CURRENT_TIMESTAMP,
               updated_at = CURRENT_TIMESTAMP
           WHERE id = ?""",
        (status, admin["id"], report_id),
    )
    db.commit()
    record_audit(
        request,
        admin,
        "report.review",
        target_type="report",
        target_id=report_id,
        summary=f"报告复核状态更新为 {status}",
    )
    return {"report": get_report_detail(report_id)}


@app.get("/api/admin/agents/{agent_name}")
def agent_detail(agent_name: str, _admin: dict = Depends(require_permission("manageModelConfig"))):
    return get_agent_detail(agent_name)


@app.get("/api/admin/users")
def list_admin_users(_admin: dict = Depends(require_super_admin)):
    return {"admins": list_admin_accounts()}


def _validate_permissions(values: Any) -> list[str]:
    if not values:
        return []
    if not isinstance(values, list):
        raise error(400, "权限列表格式不正确。")
    result: list[str] = []
    for item in values:
        key = str(item or "").strip()
        if key not in ALL_PERMISSIONS:
            raise error(400, f"包含未知权限点：{key}")
        if key not in result:
            result.append(key)
    return result


def _validate_scope(values: Any) -> list[dict[str, str]]:
    if not values:
        return []
    if not isinstance(values, list) or len(values) > 500:
        raise error(400, "数据范围格式不正确。")
    normalized: list[dict[str, str]] = []
    for entry in values:
        if not isinstance(entry, dict):
            raise error(400, "数据范围格式不正确。")
        college = str(entry.get("college") or "").strip()
        program = str(entry.get("program") or "").strip()
        class_id = str(entry.get("class") or "").strip()
        if not college or not one("SELECT id FROM campus_colleges WHERE id = ?", (college,)):
            raise error(400, "数据范围包含不存在的学院。")
        if program and not one("SELECT id FROM campus_programs WHERE id = ?", (program,)):
            raise error(400, "数据范围包含不存在的专业。")
        if class_id and not one("SELECT id FROM campus_classes WHERE id = ?", (class_id,)):
            raise error(400, "数据范围包含不存在的班级。")
        normalized.append({"college": college, "program": program, "class": class_id})
    return normalized


def _merge_scope(scope_a: list[dict[str, str]], scope_b: list[dict[str, str]]) -> list[dict[str, str]]:
    seen: set[tuple[str, str, str]] = set()
    result: list[dict[str, str]] = []
    for entry in list(scope_a) + list(scope_b):
        key = (entry.get("college") or "", entry.get("program") or "", entry.get("class") or "")
        if key in seen:
            continue
        seen.add(key)
        result.append({"college": entry.get("college") or "", "program": entry.get("program") or "", "class": entry.get("class") or ""})
    return result


@app.post("/api/admin/users", status_code=201)
async def create_admin_user(request: Request, admin: dict = Depends(require_super_admin)):
    validate_admin_origin(request)
    body = await request.json()
    email = normalize_email(body.get("email"))
    password = str(body.get("password") or "")
    name = str(body.get("name") or "").strip()
    role = str(body.get("role") or "管理员").strip()[:80] or "管理员"
    permissions = _validate_permissions(body.get("permissions"))
    scope = _validate_scope(body.get("studentScope") or body.get("student_scope"))
    if not is_valid_email(email) or len(password) < 12 or not name or len(name) > 80:
        raise error(400, "管理员邮箱、姓名或密码不符合要求；密码至少 12 位。")
    if email == MASTER_ADMIN_EMAIL:
        raise error(400, "主管理员账号由代码定义，不能重复创建。")
    if one("SELECT id FROM admin_users WHERE email = ?", (email,)):
        raise error(409, "该管理员邮箱已经存在。")
    admin_id = str(uuid4())
    db.execute(
        """
        INSERT INTO admin_users (id, email, password_hash, name, role, status, permissions, student_scope)
        VALUES (?, ?, ?, ?, ?, 'normal', ?, ?)
        """,
        (admin_id, email, hash_password(password), name, role, json.dumps(permissions), json.dumps(scope) if scope else None),
    )
    db.commit()
    record_audit(
        request,
        admin,
        "admin_user.create",
        target_type="admin_user",
        target_id=admin_id,
        summary=f"创建管理员 {email}，角色 {role}，权限 {','.join(permissions) or '无'}",
    )
    created = one("SELECT id, email, name, role, status, permissions, student_scope, last_login_at FROM admin_users WHERE id = ?", (admin_id,))
    return {"admin": sanitize_admin(created)}


@app.patch("/api/admin/users/{admin_id}")
async def update_admin_user(request: Request, admin_id: str, admin: dict = Depends(require_super_admin)):
    validate_admin_origin(request)
    target = one("SELECT id, email, name, role, status, permissions, student_scope FROM admin_users WHERE id = ?", (admin_id,))
    if not target:
        raise error(404, "管理员不存在。")
    if str(target.get("email") or "").lower() == MASTER_ADMIN_EMAIL:
        raise error(400, "主管理员账号由代码定义，不能修改或禁用。")
    body = await request.json()
    status = str(body.get("status") or target["status"]).strip()
    name = str(body.get("name") or target["name"]).strip()
    role = str(body.get("role") or target.get("role") or "管理员").strip()[:80] or "管理员"
    if status not in {"normal", "disabled"} or not name or len(name) > 80:
        raise error(400, "管理员姓名或状态不合法。")
    if "permissions" in body:
        permissions = _validate_permissions(body.get("permissions"))
    else:
        permissions = sorted(admin_permissions(target))
    if "studentScope" in body or "student_scope" in body:
        scope = _validate_scope(body.get("studentScope") or body.get("student_scope"))
    else:
        scope = admin_student_scope(target) or []
    if admin_id == admin["id"] and status != "normal":
        raise error(400, "不能禁用当前登录的管理员账号。")
    db.execute(
        "UPDATE admin_users SET name = ?, role = ?, status = ?, permissions = ?, student_scope = ? WHERE id = ?",
        (name, role, status, json.dumps(permissions), json.dumps(scope) if scope else None, admin_id),
    )
    if status == "disabled":
        db.execute("DELETE FROM admin_sessions WHERE admin_user_id = ?", (admin_id,))
    db.commit()
    record_audit(
        request,
        admin,
        "admin_user.update",
        target_type="admin_user",
        target_id=admin_id,
        summary=f"更新管理员 {target['email']}：status={status}，权限 {','.join(permissions) or '无'}",
    )
    updated = one("SELECT id, email, name, role, status, permissions, student_scope, last_login_at FROM admin_users WHERE id = ?", (admin_id,))
    return {"admin": sanitize_admin(updated)}


def _permission_request_payload(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row["id"],
        "requesterEmail": row["requester_email"],
        "requesterName": row.get("requester_name") or "",
        "permissions": json.loads(row.get("permissions") or "[]"),
        "studentScope": json.loads(row.get("student_scope") or "[]"),
        "reason": row.get("reason") or "",
        "status": row.get("status"),
        "reviewedBy": row.get("reviewed_by"),
        "reviewedAt": str(row.get("reviewed_at") or "-"),
        "createdAt": str(row.get("created_at") or "-"),
    }


def get_permission_request(request_id: str) -> dict[str, Any] | None:
    row = one("SELECT * FROM admin_permission_requests WHERE id = ?", (request_id,))
    return _permission_request_payload(row) if row else None


@app.post("/api/admin/permission-requests", status_code=201)
async def create_permission_request(request: Request, admin: dict = Depends(require_admin)):
    validate_admin_origin(request)
    body = await request.json()
    permissions = _validate_permissions(body.get("permissions"))
    scope = _validate_scope(body.get("studentScope") or body.get("student_scope"))
    reason = str(body.get("reason") or "").strip()[:500]
    if not permissions and not scope:
        raise error(400, "请选择需要申请的权限或数据范围。")
    if one(
        "SELECT id FROM admin_permission_requests WHERE admin_user_id = ? AND status = 'pending'",
        (admin["id"],),
    ):
        raise error(409, "已存在待审核的权限申请，请等待超级管理员处理。")
    req_id = str(uuid4())
    db.execute(
        """
        INSERT INTO admin_permission_requests
        (id, admin_user_id, requester_email, requester_name, permissions, student_scope, reason, status)
        VALUES (?, ?, ?, ?, ?, ?, ?, 'pending')
        """,
        (req_id, admin["id"], admin["email"], admin.get("name") or "", json.dumps(permissions), json.dumps(scope) if scope else None, reason or None),
    )
    db.commit()
    record_audit(
        request,
        admin,
        "permission_request.create",
        target_type="permission_request",
        target_id=req_id,
        summary=f"申请权限 {','.join(permissions) or '无'} 与数据范围",
    )
    return {"request": get_permission_request(req_id)}


@app.get("/api/admin/permission-requests")
def list_permission_requests(admin: dict = Depends(require_admin)):
    if is_super_admin(admin):
        rows = all_rows("SELECT * FROM admin_permission_requests ORDER BY created_at DESC LIMIT 200")
    else:
        rows = all_rows(
            "SELECT * FROM admin_permission_requests WHERE admin_user_id = ? ORDER BY created_at DESC LIMIT 100",
            (admin["id"],),
        )
    return {"requests": [_permission_request_payload(row) for row in rows]}


@app.patch("/api/admin/permission-requests/{request_id}")
async def review_permission_request(request: Request, request_id: str, admin: dict = Depends(require_super_admin)):
    validate_admin_origin(request)
    body = await request.json()
    action = str(body.get("action") or "").strip()
    if action not in {"approve", "reject"}:
        raise error(400, "审核动作必须是 approve 或 reject。")
    target = one("SELECT * FROM admin_permission_requests WHERE id = ?", (request_id,))
    if not target:
        raise error(404, "权限申请不存在。")
    if target.get("status") != "pending":
        raise error(400, "该申请已处理，不能重复审核。")
    if action == "approve":
        current = one("SELECT permissions, student_scope FROM admin_users WHERE id = ?", (target["admin_user_id"],))
        if current:
            merged_perms = sorted(set(admin_permissions(current)) | set(json.loads(target["permissions"] or "[]")))
            merged_scope = _merge_scope(admin_student_scope(current) or [], json.loads(target["student_scope"] or "[]"))
            db.execute(
                "UPDATE admin_users SET permissions = ?, student_scope = ? WHERE id = ?",
                (json.dumps(merged_perms), json.dumps(merged_scope) if merged_scope else None, target["admin_user_id"]),
            )
        db.execute(
            "UPDATE admin_permission_requests SET status = 'approved', reviewed_by = ?, reviewed_at = CURRENT_TIMESTAMP WHERE id = ?",
            (admin["email"], request_id),
        )
        db.commit()
        record_audit(request, admin, "permission_request.approve", target_type="permission_request", target_id=request_id, summary=f"通过 {target['requester_email']} 的权限申请")
        return {"request": get_permission_request(request_id)}
    db.execute(
        "UPDATE admin_permission_requests SET status = 'rejected', reviewed_by = ?, reviewed_at = CURRENT_TIMESTAMP WHERE id = ?",
        (admin["email"], request_id),
    )
    db.commit()
    record_audit(request, admin, "permission_request.reject", target_type="permission_request", target_id=request_id, summary=f"驳回 {target['requester_email']} 的权限申请")
    return {"request": get_permission_request(request_id)}


@app.patch("/api/admin/settings")
async def update_settings(request: Request, admin: dict = Depends(require_super_admin)):
    validate_admin_origin(request)
    body = await request.json()
    allowed_fields = {
        "OPENAI_REALTIME_API_KEY": "openaiApiKey",
        "OPENAI_REALTIME_MODEL": "openaiRealtimeModel",
        "OPENAI_REALTIME_VOICE": "openaiVoice",
        "OPENAI_API_KEY": "reportOpenaiApiKey",
        "OPENAI_FOLLOWUP_MODEL": "reportOpenaiModel",
        "QWEN_API_KEY": "reportQwenApiKey",
        "QWEN_MODEL": "reportQwenModel",
        "REPORT_PROVIDER_ORDER": "reportProviderOrder",
        "AI_REPORT_TIMEOUT": "reportTimeout",
        "AI_REPORT_HTTP_RETRIES": "reportRetries",
        "DASHSCOPE_API_KEY": "dashscopeApiKey",
        "DASHSCOPE_TTS_MODEL": "qwenTtsModel",
        "DASHSCOPE_TTS_VOICE": "qwenTtsVoice",
        "DASHSCOPE_TTS_REGION": "qwenTtsRegion",
        "DASHSCOPE_WORKSPACE_ID": "qwenTtsWorkspaceId",
        "QWEN_OMNI_REALTIME_MODEL": "qwenOmniModel",
        "QWEN_OMNI_REALTIME_VOICE": "qwenOmniVoice",
        "QWEN_OMNI_REALTIME_REGION": "qwenOmniRegion",
        "QWEN_OMNI_REALTIME_WORKSPACE_ID": "qwenOmniWorkspaceId",
        "QWEN_OMNI_REALTIME_WEBRTC_ENDPOINT": "qwenOmniEndpoint",
    }
    updates: dict[str, str] = {}
    for env_key, body_key in allowed_fields.items():
        value = str(body.get(body_key) or "").strip()
        if value:
            if any(character in value for character in ("\r", "\n", "\0")):
                raise error(400, f"配置项 {body_key} 包含非法字符。")
            updates[env_key] = value
    for key in ("OPENAI_API_KEY", "QWEN_API_KEY"):
        if key in updates and (len(updates[key]) > 4096 or is_placeholder_secret(updates[key])):
            raise error(400, "报告 API Key 格式无效，请填写真实密钥。")
    for key in ("OPENAI_FOLLOWUP_MODEL", "QWEN_MODEL"):
        if key in updates and (len(updates[key]) > 120 or not all(character.isalnum() or character in "-._:/" for character in updates[key])):
            raise error(400, "报告模型名称格式无效。")
    if "REPORT_PROVIDER_ORDER" in updates:
        providers = [item.strip().lower() for item in updates["REPORT_PROVIDER_ORDER"].split(",") if item.strip()]
        allowed_providers = {"openai", "qwen", "kimi", "deepseek", "custom"}
        if not providers or len(providers) != len(set(providers)) or any(item not in allowed_providers for item in providers):
            raise error(400, "报告供应商顺序无效。")
        updates["REPORT_PROVIDER_ORDER"] = ",".join(providers)
    for key, minimum, maximum in (("AI_REPORT_TIMEOUT", 10, 180), ("AI_REPORT_HTTP_RETRIES", 0, 4)):
        if key not in updates:
            continue
        try:
            number = int(updates[key])
        except ValueError as exc:
            raise error(400, "报告超时和重试次数必须是整数。") from exc
        if number < minimum or number > maximum:
            raise error(400, f"配置项 {key} 必须在 {minimum}-{maximum} 之间。")
        updates[key] = str(number)
    if not updates:
        raise error(400, "请至少填写一个要保存的配置。")
    write_env_values(USER_BACKEND_ENV_PATH, updates)
    for key, value in updates.items():
        os.environ[key] = value
    record_audit(
        request,
        admin,
        "settings.update",
        target_type="system_settings",
        target_id="ai-providers",
        summary="更新配置项：" + ", ".join(sorted(updates)),
    )
    return {"settings": build_settings(), "message": "配置已保存。复盘报告文本模型会在下一次请求时热加载；实时语音配置需重启用户端后端。"}


@app.post("/api/admin/settings/report-providers/test")
def test_report_providers(request: Request, admin: dict = Depends(require_super_admin)):
    validate_admin_origin(request)
    from backend.src.provider_diagnostics import diagnose_report_providers

    results = diagnose_report_providers(make_request=True)
    record_audit(
        request,
        admin,
        "settings.report_providers.test",
        target_type="system_settings",
        target_id="report-providers",
        summary="测试复盘报告供应商：" + ", ".join(f"{item['provider']}={item['status']}" for item in results),
    )
    return {
        "providers": results,
        "ok": any(item.get("status") == "ok" for item in results),
    }
