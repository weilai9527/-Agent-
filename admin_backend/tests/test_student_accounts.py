from __future__ import annotations

import os
from pathlib import Path
import subprocess
import sys
import textwrap

from cryptography.fernet import Fernet


def test_admin_imports_reveals_exports_and_clears_temporary_passwords(tmp_path):
    project_dir = Path(__file__).resolve().parents[2]
    database_path = tmp_path / "student-accounts.sqlite"
    environment = os.environ.copy()
    environment.update(
        {
            "APP_ENV": "production",
            "DB_ENGINE": "sqlite",
            "SQLITE_PATH": str(database_path),
            "ADMIN_BOOTSTRAP_EMAIL": "root@example.com",
            "ADMIN_BOOTSTRAP_PASSWORD": "RootAdminPassword123!",
            "ADMIN_MASTER_EMAIL": "root@example.com",
            "ADMIN_MASTER_PASSWORD": "RootAdminPassword123!",
            "ADMIN_BOOTSTRAP_ROLE": "super_admin",
            "ADMIN_COOKIE_SECURE": "false",
            "STUDENT_TEMP_PASSWORD_KEY": Fernet.generate_key().decode("ascii"),
        }
    )
    script = textwrap.dedent(
        '''
        from io import BytesIO
        from fastapi.testclient import TestClient
        from openpyxl import Workbook, load_workbook

        import backend.src.database
        import backend.src.main as candidate_main
        import admin_backend.src.main as admin_main

        origin = {"Origin": "http://127.0.0.1:5174"}
        admin = TestClient(admin_main.app)
        login = admin.post(
            "/api/admin/auth/login",
            headers=origin,
            json={"email": "root@example.com", "password": "RootAdminPassword123!"},
        )
        assert login.status_code == 200, login.text

        workbook = Workbook()
        sheet = workbook.active
        counselor = "\\u738b\\u8001\\u5e08"
        sheet.append(["\\u5b66\\u9662", "\\u5b66\\u53f7", "\\u59d3\\u540d", "\\u6027\\u522b", "\\u73ed\\u7ea7", "\\u8f85\\u5bfc\\u5458", "\\u8d26\\u53f7\\u72b6\\u6001", "\\u5b66\\u751f\\u72b6\\u6001", "\\u6ce8\\u518c\\u65f6\\u95f4", "\\u4fee\\u6539\\u65f6\\u95f4"])
        sheet.append(["\\u8ba1\\u7b97\\u673a\\u5b66\\u9662", "20260001", "\\u5f20\\u4e09", "\\u7537", "\\u8f6f\\u4ef6\\u4e00\\u73ed", counselor, "\\u6b63\\u5e38", "\\u5728\\u6821", "2026-09-01", "2026-09-01"])
        sheet.append(["\\u8ba1\\u7b97\\u673a\\u5b66\\u9662", "20260002", "\\u674e\\u56db", "\\u5973", "\\u8f6f\\u4ef6\\u4e00\\u73ed", counselor, "\\u6b63\\u5e38", "\\u5728\\u6821", "2026-09-01", "2026-09-01"])
        content = BytesIO()
        workbook.save(content)

        imported = admin.post(
            "/api/admin/student-accounts/import",
            headers=origin,
            files={"file": ("students.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert imported.status_code == 200, imported.text
        assert imported.json()["result"]["created"] == 2

        students = admin_main.all_rows(
            "SELECT id, student_no, password_hash, temp_password_encrypted FROM users ORDER BY student_no"
        )
        assert len(students) == 2
        assert students[0]["password_hash"] != students[1]["password_hash"]
        assert "20260001" not in students[0]["temp_password_encrypted"]

        revealed = []
        for student in students:
            response = admin.get(f"/api/admin/student-accounts/{student['id']}/temporary-password")
            assert response.status_code == 200, response.text
            revealed.append(response.json()["temporaryPassword"])
        assert revealed[0] != revealed[1]

        exported = admin.get(f"/api/admin/student-accounts/export?counselor={counselor}&class_name=\\u8f6f\\u4ef6\\u4e00\\u73ed&admission_year=2026")
        assert exported.status_code == 200, exported.text
        exported_workbook = load_workbook(BytesIO(exported.content), data_only=True)
        exported_rows = list(exported_workbook.active.iter_rows(values_only=True))
        assert len(exported_rows) == 3
        assert exported_rows[1][1] == "20260001"
        assert exported_rows[1][2] == "2026"
        assert exported_rows[1][7] == revealed[0]

        empty_year = admin.get("/api/admin/student-accounts/export?admission_year=2025")
        assert empty_year.status_code == 200
        empty_workbook = load_workbook(BytesIO(empty_year.content), data_only=True)
        assert empty_workbook.active.max_row == 1

        student_client = TestClient(candidate_main.app)
        student_login = student_client.post(
            "/api/auth/login",
            json={"student_no": "20260001", "password": revealed[0]},
        )
        assert student_login.status_code == 200, student_login.text
        assert student_login.json()["user"]["mustChangePassword"] is True
        changed = student_client.post(
            "/api/auth/change-initial-password",
            json={"password": "StudentPassword123!", "confirm_password": "StudentPassword123!"},
        )
        assert changed.status_code == 200, changed.text

        hidden = admin.get(f"/api/admin/student-accounts/{students[0]['id']}/temporary-password")
        assert hidden.status_code == 409, hidden.text
        cleared = admin_main.one(
            "SELECT temp_password_encrypted, must_change_password FROM users WHERE id = ?",
            (students[0]["id"],),
        )
        assert cleared["temp_password_encrypted"] is None
        assert cleared["must_change_password"] == 0

        remaining_export = admin.get("/api/admin/student-accounts/export")
        remaining_workbook = load_workbook(BytesIO(remaining_export.content), data_only=True)
        assert remaining_workbook.active.max_row == 2

        # Re-importing a legacy activated student adopts the existing user ID,
        # keeping their profile and interview history attached to the account.
        admin_main.db.execute(
            "INSERT INTO users (id, email, password_hash, name) VALUES (?, ?, ?, ?)",
            ("legacy-student", "20260003@student.local", "legacy-hash", "Legacy Student"),
        )
        admin_main.db.execute(
            "INSERT INTO profiles (id, user_id, nickname) VALUES (?, ?, ?)",
            ("legacy-profile", "legacy-student", "Legacy Student"),
        )
        admin_main.db.execute(
            "INSERT INTO student_registrations (id, student_no, name, user_id) VALUES (?, ?, ?, ?)",
            ("legacy-registration", "20260003", "Legacy Student", "legacy-student"),
        )
        admin_main.db.commit()
        legacy_workbook = Workbook()
        legacy_sheet = legacy_workbook.active
        legacy_sheet.append(["\\u5b66\\u53f7", "\\u59d3\\u540d"])
        legacy_sheet.append(["20260003", "Legacy Student"])
        legacy_content = BytesIO()
        legacy_workbook.save(legacy_content)
        migrated = admin.post(
            "/api/admin/student-accounts/import",
            headers=origin,
            files={"file": ("legacy.xlsx", legacy_content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert migrated.status_code == 200, migrated.text
        assert migrated.json()["result"]["migrated"] == 1
        assert migrated.json()["result"]["created"] == 0
        legacy = admin_main.one("SELECT id, student_no, must_change_password FROM users WHERE id = ?", ("legacy-student",))
        assert legacy == {"id": "legacy-student", "student_no": "20260003", "must_change_password": 1}
        assert admin_main.one("SELECT id FROM profiles WHERE user_id = ?", ("legacy-student",))["id"] == "legacy-profile"
        assert admin.get("/api/admin/student-accounts/legacy-student/temporary-password").status_code == 200

        # A deleted legacy registration may still leave its synthetic-email user.
        admin_main.db.execute(
            "INSERT INTO users (id, email, password_hash, name) VALUES (?, ?, ?, ?)",
            ("orphan-student", "20260004@student.local", "legacy-hash", "Orphan Student"),
        )
        admin_main.db.commit()
        orphan_workbook = Workbook()
        orphan_sheet = orphan_workbook.active
        orphan_sheet.append(["\\u5b66\\u53f7", "\\u59d3\\u540d"])
        orphan_sheet.append(["20260004", "Orphan Student"])
        orphan_content = BytesIO()
        orphan_workbook.save(orphan_content)
        adopted = admin.post(
            "/api/admin/student-accounts/import",
            headers=origin,
            files={"file": ("orphan.xlsx", orphan_content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert adopted.status_code == 200, adopted.text
        assert adopted.json()["result"]["migrated"] == 1
        assert admin_main.one("SELECT student_no FROM users WHERE id = ?", ("orphan-student",))["student_no"] == "20260004"
        '''
    )
    completed = subprocess.run(
        [sys.executable, "-c", script],
        cwd=project_dir,
        env=environment,
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr
